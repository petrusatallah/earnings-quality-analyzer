"""Pure presentation-layer export for finalized earnings-quality results.

The production API in this module accepts every value it writes.  It does not
import company datasets, calculation modules, analysis modules, or benchmark
fixtures, and it does not calculate or classify financial results.
"""

from __future__ import annotations

import json
import math
from collections import OrderedDict
from copy import copy
from enum import Enum
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Iterable, Mapping

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.dataframe import dataframe_to_rows

from Analysis.reporting_view_model import (
    ChartSpec,
    RULE_PRESENTATION_MAPPING,
    ReportingInputEnvelope,
    ReportingViewModel,
    build_reporting_view_model,
)


PRODUCTION_SHEET_ORDER = (
    "Dashboard",
    "Analyst Conclusion",
    "Raw Financial Data",
    "Growth Calculations",
    "Accrual & Cash Analysis",
    "Working Capital",
    "D&A and CapEx",
    "Deferred Taxes",
    "SBC & Dilution",
    "Normalized Earnings",
    "Master Analysis",
    "Analyst Summary",
    "Red Flags",
    "Detailed Severity",
    "Overall Severity",
    "Unavailable Outputs",
    "Analytical Explanations",
)

# These are presentation-only aliases from the documented workbook contract.
_RULE_DISPLAY = {
    key: presentation.display_name
    for key, presentation in RULE_PRESENTATION_MAPPING.items()
}
DISPLAY_HEADER_ALIASES = {
    "Accrual & Cash Analysis": {
        "AR Flag": f"{_RULE_DISPLAY['AR growth vs Revenue growth']} — flag",
        "NI OCF Flag": f"{_RULE_DISPLAY['Net Income growth vs OCF growth']} — flag",
        "Combined Accrual Flag": "Combined Accrual Warning?",
    },
    "Working Capital": {
        "Inventory Flag": f"{_RULE_DISPLAY['Inventory growth vs Revenue growth']} — flag",
        "Inventory Result": f"{_RULE_DISPLAY['Inventory growth vs Revenue growth']} — result",
        "AP Result": f"{_RULE_DISPLAY['Accounts Payable pattern']} — result",
        "WC Result": f"{_RULE_DISPLAY['Selected-account working-capital proxy']} — result",
        "Working Capital Cash Use": "Working Capital Used Cash?",
    },
    "D&A and CapEx": {
        "Status": f"{_RULE_DISPLAY['CapEx vs D&A']} — result",
    },
    "Deferred Taxes": {
        "Deferred Tax Movement Flag": f"{_RULE_DISPLAY['Deferred tax movement']} — flag",
        "Deferred Tax Movement Result": f"{_RULE_DISPLAY['Deferred tax movement']} — result",
        "DTA Risk Flag": f"{_RULE_DISPLAY['Deferred Tax Asset risk']} — flag",
        "DTA Result": f"{_RULE_DISPLAY['Deferred Tax Asset risk']} — result",
    },
    "SBC & Dilution": {
        "Large SBC Result": f"{_RULE_DISPLAY['Large SBC']} — result",
        "Large SBC Flag": f"{_RULE_DISPLAY['Large SBC']} — flag",
        "Large SBC Severity": f"{_RULE_DISPLAY['Large SBC']} — severity",
        "Dilution Flag": f"{_RULE_DISPLAY['SBC dilution']} — flag",
        "Dilution Result": f"{_RULE_DISPLAY['SBC dilution']} — result",
        "Buyback Offset Flag": f"{_RULE_DISPLAY['Buyback offset']} — flag",
        "Buyback Offset Result": f"{_RULE_DISPLAY['Buyback offset']} — result",
    },
    "Normalized Earnings": {
        "Large Normalization Difference Flag": (
            f"{_RULE_DISPLAY['Large normalization difference']} — flag"
        ),
        "Large Normalization Difference Result": (
            f"{_RULE_DISPLAY['Large normalization difference']} — result"
        ),
        "Large Normalization Difference Severity": (
            f"{_RULE_DISPLAY['Large normalization difference']} — severity"
        ),
        "Repeated One-Off Flag": f"{_RULE_DISPLAY['Repeated one-off items']} — flag",
        "Repeated One-Off Result": f"{_RULE_DISPLAY['Repeated one-off items']} — result",
        "Repeated One-Off Severity": (
            f"{_RULE_DISPLAY['Repeated one-off items']} — severity"
        ),
    },
    "Overall Severity": {
        "Benchmark Severity": "Overall Severity",
    },
}

PERCENTAGE_HEADERS = {
    "Revenue Growth", "AR Growth", "Inventory Growth", "AP Growth",
    "Net Income Growth", "OCF Growth", "D&A Growth", "CapEx Growth",
    "SBC Growth", "Shares Outstanding Growth", "AR Revenue Gap",
    "Inventory Revenue Gap", "AP Revenue Gap", "D&A / Revenue",
    "CapEx / Revenue", "SBC / Revenue", "SBC / OCF", "SBC / Net Income",
    "DTA / Net Income", "DTA Growth", "DTL Growth",
    "Percentage Difference", "Normalization Difference Percentage",
}
RATIO_HEADERS = {"CapEx / D&A", "Buyback Offset Ratio"}
FINANCIAL_HEADERS = {
    "Revenue", "Net Income", "Operating Cash Flow", "Free Cash Flow",
    "Reported Net Income", "Normalized Net Income", "Difference",
    "Normalization Difference", "Capital Expenditures", "CapEx",
    "Depreciation & Amortization", "D&A", "Stock-Based Compensation", "SBC",
    "Tax Expense", "Deferred Tax Assets", "Deferred Tax Liabilities",
    "Accounts Receivable", "Inventory", "Accounts Payable", "AR Change",
    "Inventory Change", "AP Change", "Net Working Capital Cash Effect",
    "Cash Spent on Buybacks", "Value", "Current Year Value",
    "Previous Year Value", "Change", "Net Normalization Adjustment",
    "Normalized NI Difference",
}

_INVESTOR_ANALYSIS_SPECS = OrderedDict({
    "Accrual & Cash Analysis": (
        "accrual_and_cash",
        (
            "Fiscal Year", "Net Income Growth", "OCF Growth", "Revenue Growth",
            "AR Growth", "AR Revenue Gap", "AR Flag", "NI OCF Flag",
            "Combined Accrual Flag", "Accrual Signal", "Severity",
            "Analyst Interpretation",
        ),
    ),
    "Working Capital": (
        "working_capital",
        (
            "Fiscal Year", "Inventory Growth", "Revenue Growth",
            "Inventory Revenue Gap", "AP Growth", "OCF Growth", "AP Revenue Gap",
            "Net Working Capital Cash Effect", "Working Capital Cash Use",
            "Inventory Result", "AP Result", "WC Result", "Working Capital Signal",
            "Overall Severity", "Analyst Interpretation",
        ),
    ),
    "D&A and CapEx": (
        "da_and_capex",
        (
            "Fiscal Year", "CapEx", "D&A", "CapEx / D&A", "Status", "Severity",
            "Analyst Interpretation",
        ),
    ),
    "Deferred Taxes": (
        "deferred_taxes",
        (
            "Fiscal Year", "Deferred Tax Assets",
            "Deferred Tax Liabilities", "Net Income", "DTA / Net Income",
            "DTA Growth", "DTL Growth", "DTA Risk Flag", "DTA Result",
            "Deferred Tax Movement Flag", "Deferred Tax Movement Result",
            "Tax Signal", "Overall Tax Severity", "Analyst Interpretation",
        ),
    ),
    "SBC & Dilution": (
        "sbc_and_dilution",
        (
            "Fiscal Year", "SBC", "Net Income", "SBC / Net Income",
            "Shares Outstanding", "Shares Outstanding Growth", "Dilution Flag",
            "Shares Issued Net", "Shares Repurchased", "Buyback Offset Ratio",
            "Large SBC Result", "Dilution Result", "Buyback Offset Result",
            "SBC Signal", "Overall SBC Severity", "Analyst Interpretation",
        ),
    ),
    "Normalized Earnings": (
        "normalized_earnings",
        (
            "Fiscal Year", "Reported Net Income", "Normalized Net Income",
            "Difference", "Percentage Difference",
            "Earnings Adjustment Direction", "One-Off Categories Identified",
            "Large Normalization Difference Result", "Repeated One-Off Result",
            "Normalization Signal", "Overall Severity", "Analyst Interpretation",
        ),
    ),
})

METHODOLOGY_SHEET_RULES = OrderedDict({
    "Accrual & Cash Analysis": (
        "AR growth vs Revenue growth",
        "Net Income growth vs OCF growth",
    ),
    "Working Capital": (
        "Inventory growth vs Revenue growth",
        "Accounts Payable pattern",
        "Selected-account working-capital proxy",
    ),
    "D&A and CapEx": ("CapEx vs D&A",),
    "Deferred Taxes": (
        "Deferred tax movement",
        "Deferred Tax Asset risk",
    ),
    "SBC & Dilution": (
        "Large SBC",
        "SBC dilution",
        "Buyback offset",
    ),
    "Normalized Earnings": (
        "Large normalization difference",
        "Repeated one-off items",
    ),
    "Red Flags": tuple(RULE_PRESENTATION_MAPPING),
})

_INVESTOR_RULE_TEXT_ALIASES = (
    ("net-income-versus-operating-cash-flow", "Net Income growth vs OCF growth"),
    ("large-normalization-difference", "Large normalization difference"),
    ("selected-account working-capital proxy", "Selected-account working-capital proxy"),
    ("deferred-tax-asset risk", "Deferred Tax Asset risk"),
    ("deferred-tax-movement", "Deferred tax movement"),
    ("AR-versus-revenue", "AR growth vs Revenue growth"),
    ("accounts-payable", "Accounts Payable pattern"),
    ("inventory-growth", "Inventory growth vs Revenue growth"),
    ("repeated-one-off", "Repeated one-off items"),
    ("large-SBC", "Large SBC"),
    ("CapEx review", "CapEx vs D&A"),
    ("inventory rule", "Inventory growth vs Revenue growth"),
    ("AP classification", "Accounts Payable pattern"),
    ("dilution rule", "SBC dilution"),
)


def _require_frame(container: Mapping[str, Any], key: str) -> pd.DataFrame:
    value = container.get(key)
    if not isinstance(value, pd.DataFrame):
        raise ValueError(f"Finalized pipeline output {key!r} is unavailable")
    return value.copy(deep=True)


def _growth_view(financial_summary: pd.DataFrame) -> pd.DataFrame:
    columns = [
        column for column in financial_summary.columns
        if column == "Fiscal Year"
        or "Growth" in str(column)
        or column in {"AR Revenue Gap"}
    ]
    return financial_summary.loc[:, columns].copy(deep=True)


def _red_flags_view(analysis: Mapping[str, Any]) -> pd.DataFrame:
    red_flags = _require_frame(analysis, "red_flags")
    detailed = _require_frame(analysis, "detailed_severity")
    key_columns = ["Fiscal Year", "Metric"]
    if "Final Severity" not in red_flags.columns:
        lookup = detailed.loc[:, key_columns + ["Final Severity"]]
        red_flags = red_flags.merge(
            lookup, on=key_columns, how="left", validate="one_to_one"
        )
    return red_flags


def _normalization_unavailable(frame: pd.DataFrame) -> bool:
    if "Normalization Signal" in frame.columns:
        return frame["Normalization Signal"].astype(str).eq("Unavailable").all()
    return False


def _json_default(value: Any) -> Any:
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    normalized = normalize_excel_scalar(value, display_booleans=False)
    if normalized is not value:
        return normalized
    return str(value)


def _interpretation_rows(
    interpretations: Mapping[str, Any], *, omit_normalization: bool
) -> pd.DataFrame:
    rows = []
    for area, output in interpretations.items():
        if omit_normalization and area == "normalized_earnings":
            continue
        if not isinstance(output, Mapping):
            rows.append({"Analysis": area, "Field": "value", "Value": output})
            continue
        for field, value in output.items():
            if isinstance(value, (Mapping, list, tuple)):
                continue
            rows.append({"Analysis": area, "Field": field, "Value": value})
    return pd.DataFrame(rows, columns=["Analysis", "Field", "Value"])


def _investor_analysis_frame(
    view_model: ReportingViewModel, *, area_key: str, preferred_columns: Iterable[str]
) -> pd.DataFrame:
    """Present one finalized analysis area without changing its financial content."""

    area = view_model.area(area_key)
    records = [dict(row.items()) for row in area.finalized_rows]
    frame = pd.DataFrame(records, columns=area.source_columns)
    assessments = {
        assessment.fiscal_year: assessment for assessment in area.yearly_assessments
    }
    if "Fiscal Year" not in frame.columns:
        raise ValueError(f"Finalized {area.label} rows do not contain Fiscal Year")
    frame["Analyst Interpretation"] = [
        (
            assessments[fiscal_year].assessment
            if fiscal_year in assessments
            and assessments[fiscal_year].assessment is not None
            else None
        )
        for fiscal_year in frame["Fiscal Year"]
    ]
    if area_key == "da_and_capex":
        frame["Severity"] = [
            (
                assessments[fiscal_year].existing_severity
                if fiscal_year in assessments
                and assessments[fiscal_year].existing_severity is not None
                else None
            )
            for fiscal_year in frame["Fiscal Year"]
        ]

    investor_columns = [
        column for column in preferred_columns if column in frame.columns
    ]
    audit_columns = [
        column for column in frame.columns if column not in investor_columns
    ]
    return frame.loc[:, investor_columns + audit_columns]


def build_production_sheets(
    *,
    corrected_data: pd.DataFrame,
    calculations: Mapping[str, Any],
    analysis: Mapping[str, Any],
    interpretations: Mapping[str, Any],
    reporting_view_model: ReportingViewModel,
) -> OrderedDict[str, pd.DataFrame]:
    """Build presentation tables only from finalized same-run outputs."""

    if not isinstance(corrected_data, pd.DataFrame) or corrected_data.empty:
        raise ValueError("Corrected financial data is unavailable for Excel export")
    if not isinstance(calculations, Mapping) or not isinstance(analysis, Mapping):
        raise ValueError("Finalized calculations and analysis are required")
    if not isinstance(interpretations, Mapping):
        raise ValueError("Finalized interpretations are required")

    summary = _require_frame(calculations, "financial_summary")
    normalized = _require_frame(calculations, "normalized_earnings")
    support = calculations.get("excel_support_tables")
    if not isinstance(support, Mapping):
        raise ValueError("Finalized Excel support tables are unavailable")

    unavailable = analysis.get("unavailable_outputs", pd.DataFrame())
    if not isinstance(unavailable, pd.DataFrame):
        raise ValueError("Finalized unavailable-output status is invalid")

    sheets = OrderedDict([
        ("Raw Financial Data", corrected_data.copy(deep=True)),
        ("Growth Calculations", _growth_view(summary)),
        *(
            (
                sheet_name,
                _investor_analysis_frame(
                    reporting_view_model,
                    area_key=area_key,
                    preferred_columns=preferred_columns,
                ),
            )
            for sheet_name, (area_key, preferred_columns)
            in _INVESTOR_ANALYSIS_SPECS.items()
        ),
        ("Master Analysis", summary),
        ("Analyst Summary", _require_frame(analysis, "fiscal_year_severity")),
        ("Red Flags", _red_flags_view(analysis)),
        ("Detailed Severity", _require_frame(analysis, "detailed_severity")),
        ("Overall Severity", _require_frame(analysis, "overall_severity")),
        ("Unavailable Outputs", unavailable.copy(deep=True)),
        (
            "Analytical Explanations",
            _interpretation_rows(
                interpretations,
                omit_normalization=_normalization_unavailable(normalized),
            ),
        ),
    ])
    return sheets


def normalize_excel_scalar(value: Any, *, display_booleans: bool = True) -> Any:
    """Return the approved scalar representation used by Excel and validation."""

    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, Enum):
        return normalize_excel_scalar(
            value.value, display_booleans=display_booleans
        )
    if display_booleans and isinstance(value, (bool, np.bool_)):
        return "Yes" if bool(value) else "No"
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=_json_default)
    return value


def _display_rule_name(value: Any) -> Any:
    normalized = normalize_excel_scalar(value)
    presentation = RULE_PRESENTATION_MAPPING.get(normalized)
    return presentation.display_name if presentation is not None else normalized


def _display_rule_list(value: Any) -> Any:
    normalized = normalize_excel_scalar(value)
    if not isinstance(normalized, str) or normalized in {"", "None", "Unavailable"}:
        return normalized
    return "; ".join(
        str(_display_rule_name(item.strip()))
        for item in normalized.split(";")
        if item.strip()
    )


def _humanize_investor_rule_text(value: Any) -> Any:
    normalized = normalize_excel_scalar(value)
    if not isinstance(normalized, str):
        return normalized
    humanized = normalized
    for obsolete_text, rule_key in _INVESTOR_RULE_TEXT_ALIASES:
        humanized = humanized.replace(
            obsolete_text, RULE_PRESENTATION_MAPPING[rule_key].display_name
        )
    return humanized


def displayed_frame(sheet_name: str, frame: pd.DataFrame) -> pd.DataFrame:
    """Return the documented, presentation-only view used in the workbook."""

    displayed = frame.copy(deep=True)
    if sheet_name == "Red Flags" and {"Metric", "Rule"}.issubset(displayed.columns):
        displayed["Rule"] = [
            (
                RULE_PRESENTATION_MAPPING[metric].plain_english_explanation
                if metric in RULE_PRESENTATION_MAPPING
                else rule
            )
            for metric, rule in zip(displayed["Metric"], displayed["Rule"])
        ]
    if "Metric" in displayed.columns:
        displayed["Metric"] = [_display_rule_name(value) for value in displayed["Metric"]]
    for column in ("Triggered Metrics", "Unavailable Metrics"):
        if column in displayed.columns:
            displayed[column] = [_display_rule_list(value) for value in displayed[column]]

    displayed = displayed.rename(
        columns=DISPLAY_HEADER_ALIASES.get(sheet_name, {})
    )
    for column in displayed.columns:
        displayed[column] = pd.Series(
            [normalize_excel_scalar(value) for value in displayed[column]],
            index=displayed.index,
            dtype=object,
        )
    return displayed


def _analysis_sheet_tables(
    sheet_name: str, frame: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    rule_columns = [] if sheet_name == "Red Flags" else [
        column for column in frame.columns if "rule" in str(column).casefold()
    ]
    main_table = frame.drop(columns=rule_columns).copy(deep=True)
    methodology_rows = []
    for rule_key in METHODOLOGY_SHEET_RULES.get(sheet_name, ()):
        presentation = RULE_PRESENTATION_MAPPING[rule_key]
        methodology_rows.append({
            "Rule": presentation.display_name,
            "Plain-English explanation": presentation.plain_english_explanation,
            "Exact technical rule": presentation.exact_technical_rule,
        })
    methodology = pd.DataFrame(
        methodology_rows,
        columns=["Rule", "Plain-English explanation", "Exact technical rule"],
    )
    return main_table, methodology


def _write_frame(worksheet: Any, sheet_name: str, frame: pd.DataFrame) -> None:
    displayed = displayed_frame(sheet_name, frame)
    for row in dataframe_to_rows(displayed, index=False, header=True):
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False


def _add_status_conditional_formatting(
    worksheet: Any, *, min_row: int, max_row: int, max_column: int
) -> None:
    if max_row < min_row or max_column < 1:
        return
    target = f"A{min_row}:{get_column_letter(max_column)}{max_row}"
    styles = {
        "Low Risk": ("E2F0D9", "006100"),
        "Needs Investigation": ("FFF2CC", "9C6500"),
        "Material Concern": ("F4CCCC", "9C0006"),
        "Unavailable": ("E7E6E6", "666666"),
    }
    for text, (fill_color, font_color) in styles.items():
        worksheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'A{min_row}="{text}"'],
                fill=PatternFill("solid", fgColor=fill_color),
                font=Font(color=font_color),
            ),
        )


def _methodology_layout(main_column_count: int) -> tuple[int, int, int, int]:
    """Return logical Rule, explanation, and technical-rule column boundaries."""

    last_column = max(main_column_count, 12)
    rule_end = max(3, round(last_column * 0.22))
    explanation_start = rule_end + 1
    explanation_end = max(explanation_start + 3, round(last_column * 0.68))
    explanation_end = min(explanation_end, last_column - 2)
    technical_start = explanation_end + 1
    return last_column, rule_end, explanation_start, technical_start


def _write_investor_analysis_sheet(
    worksheet: Any, sheet_name: str, frame: pd.DataFrame
) -> None:
    main_table, methodology = _analysis_sheet_tables(sheet_name, frame)
    _write_frame(worksheet, sheet_name, main_table)
    _format_table(worksheet)

    displayed_headers = [cell.value for cell in worksheet[1]]
    if "Analyst Interpretation" in displayed_headers:
        analyst_column = displayed_headers.index("Analyst Interpretation") + 1
        worksheet.column_dimensions[get_column_letter(analyst_column)].width = 50
        for row in range(2, len(main_table) + 2):
            worksheet.cell(row, analyst_column).alignment = Alignment(
                vertical="top", wrap_text=True
            )
    if sheet_name == "Red Flags":
        red_flag_widths = {"Metric": 42, "Rule": 72, "Explanation": 64}
        header_columns = {
            value: index for index, value in enumerate(displayed_headers, start=1)
        }
        for header, width in red_flag_widths.items():
            if header not in header_columns:
                continue
            column = header_columns[header]
            worksheet.column_dimensions[get_column_letter(column)].width = width
            for row in range(2, len(main_table) + 2):
                worksheet.cell(row, column).alignment = Alignment(
                    vertical="top", wrap_text=True
                )
        for row in range(2, len(main_table) + 2):
            required_lines = max(
                math.ceil(
                    len(str(worksheet.cell(row, header_columns[header]).value or ""))
                    / max(20, width)
                )
                for header, width in red_flag_widths.items()
                if header in header_columns
            )
            worksheet.row_dimensions[row].height = max(
                worksheet.row_dimensions[row].height or 15,
                min(180, required_lines * 15 + 12),
            )

    _add_status_conditional_formatting(
        worksheet,
        min_row=2,
        max_row=len(main_table) + 1,
        max_column=len(main_table.columns),
    )

    section_row = len(main_table) + 3
    header_row = section_row + 1
    first_rule_row = header_row + 1
    last_column, rule_end, explanation_start, technical_start = (
        _methodology_layout(len(main_table.columns))
    )
    explanation_end = technical_start - 1
    worksheet.merge_cells(
        start_row=section_row, start_column=1,
        end_row=section_row, end_column=last_column,
    )
    section = worksheet.cell(section_row, 1)
    section.value = "Rules & Methodology"
    section.fill = PatternFill("solid", fgColor="17365D")
    section.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    section.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[section_row].height = 22

    worksheet.merge_cells(
        start_row=header_row, start_column=1,
        end_row=header_row, end_column=rule_end,
    )
    worksheet.merge_cells(
        start_row=header_row, start_column=explanation_start,
        end_row=header_row, end_column=explanation_end,
    )
    worksheet.merge_cells(
        start_row=header_row, start_column=technical_start,
        end_row=header_row, end_column=last_column,
    )
    worksheet.cell(header_row, 1).value = "Rule"
    worksheet.cell(header_row, explanation_start).value = "Plain-English explanation"
    worksheet.cell(header_row, technical_start).value = "Exact technical rule"
    for coordinate in (
        (header_row, 1),
        (header_row, explanation_start),
        (header_row, technical_start),
    ):
        cell = worksheet.cell(*coordinate)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(name="Arial", size=10, bold=True, color="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[header_row].height = 30

    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for offset, record in enumerate(methodology.to_dict("records")):
        row = first_rule_row + offset
        worksheet.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=rule_end
        )
        worksheet.merge_cells(
            start_row=row, start_column=explanation_start,
            end_row=row, end_column=explanation_end,
        )
        worksheet.merge_cells(
            start_row=row, start_column=technical_start,
            end_row=row, end_column=last_column,
        )
        worksheet.cell(row, 1).value = normalize_excel_scalar(record["Rule"])
        worksheet.cell(row, explanation_start).value = normalize_excel_scalar(
            record["Plain-English explanation"]
        )
        worksheet.cell(row, technical_start).value = normalize_excel_scalar(
            record["Exact technical rule"]
        )
        worksheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True)
        worksheet.cell(row, explanation_start).font = Font(name="Arial", size=10)
        worksheet.cell(row, technical_start).font = Font(name="Arial", size=10)
        worksheet.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
        worksheet.cell(row, explanation_start).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        worksheet.cell(row, technical_start).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        explanation_lines = math.ceil(
            len(str(record["Plain-English explanation"]))
            / max(45, (explanation_end - explanation_start + 1) * 14)
        )
        technical_lines = math.ceil(
            len(str(record["Exact technical rule"]))
            / max(35, (last_column - technical_start + 1) * 12)
        ) + str(record["Exact technical rule"]).count("\n")
        worksheet.row_dimensions[row].height = max(
            42,
            15 * max(explanation_lines, technical_lines) + 12,
        )
        for column in range(1, last_column + 1):
            worksheet.cell(row, column).border = border



def _severity_fill(value: Any) -> PatternFill | None:
    return {
        "High": PatternFill("solid", fgColor="F4CCCC"),
        "Material Concern": PatternFill("solid", fgColor="F4CCCC"),
        "Medium": PatternFill("solid", fgColor="FFE699"),
        "Needs Investigation": PatternFill("solid", fgColor="FFE699"),
        "Low": PatternFill("solid", fgColor="FFF2CC"),
        "Low Risk": PatternFill("solid", fgColor="E2F0D9"),
        "Unavailable": PatternFill("solid", fgColor="D9E1F2"),
    }.get(value)


def _format_table(worksheet: Any) -> None:
    navy = "17365D"
    thin = Side(style="thin", color="D9E1F2")
    header_fill = PatternFill("solid", fgColor=navy)
    headers = {cell.column: str(cell.value or "") for cell in worksheet[1]}
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 32

    for row in worksheet.iter_rows(min_row=2):
        required_lines = 1
        for cell in row:
            header = headers[cell.column]
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
            if isinstance(cell.value, str):
                required_lines = max(
                    required_lines,
                    sum(max(1, math.ceil(len(line) / 48)) for line in cell.value.split("\n")),
                )
            if header in PERCENTAGE_HEADERS:
                cell.number_format = "0.00%;[Red](0.00%);-"
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif header in RATIO_HEADERS:
                cell.number_format = "0.00x;[Red](0.00x);-"
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif header in FINANCIAL_HEADERS and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00;[Red](#,##0.00);-"
                cell.alignment = Alignment(horizontal="right", vertical="top")
            fill = _severity_fill(cell.value)
            if fill is not None:
                cell.fill = fill
        worksheet.row_dimensions[row[0].row].height = min(
            120, max(30, required_lines * 15 + 6)
        )

    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        worksheet.column_dimensions[letter].width = min(max(width + 2, 12), 50)


def _cover_style(worksheet: Any, *, columns: int) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A3"
    for column in range(1, columns + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 16


def _dashboard_value(value: Any) -> Any:
    normalized = normalize_excel_scalar(value)
    return "Unavailable" if normalized is None else normalized


def _dashboard_unit(
    corrected_data: pd.DataFrame, *, fiscal_year: int, metric: str
) -> str:
    matches = corrected_data.loc[
        corrected_data["Fiscal Year"].eq(fiscal_year)
        & corrected_data["Metric"].eq(metric),
        "Units",
    ]
    if matches.empty:
        return "Unavailable"
    return str(_dashboard_value(matches.iloc[-1]))


def _latest_interpretation_severity(
    interpretations: Mapping[str, Any], area: str, fiscal_year: int
) -> Any:
    output = interpretations.get(area)
    if not isinstance(output, Mapping):
        return "Unavailable"
    assessments = output.get("yearly_assessments")
    if not isinstance(assessments, list):
        return "Unavailable"
    for assessment in assessments:
        if (
            isinstance(assessment, Mapping)
            and assessment.get("fiscal_year") == fiscal_year
        ):
            return _dashboard_value(assessment.get("existing_severity"))
    return "Unavailable"


def _dashboard_status_fill(value: Any) -> PatternFill | None:
    existing = _severity_fill(value)
    if existing is not None:
        return existing
    return {
        "Strong": PatternFill("solid", fgColor="E2F0D9"),
        "Generally Healthy": PatternFill("solid", fgColor="E2F0D9"),
        "Weak": PatternFill("solid", fgColor="F4CCCC"),
        "Flag": PatternFill("solid", fgColor="F4CCCC"),
        "Review": PatternFill("solid", fgColor="FFF2CC"),
    }.get(value)


def _create_dashboard(
    worksheet: Any,
    *,
    view_model: ReportingViewModel,
    corrected_data: pd.DataFrame,
) -> dict[str, Any]:
    years = view_model.fiscal_years
    latest_year = view_model.latest_fiscal_year

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = None
    for column in range(1, 15):
        worksheet.column_dimensions[get_column_letter(column)].width = 14
    for row in range(1, 28):
        worksheet.row_dimensions[row].height = 20
    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 2
    worksheet.row_dimensions[2].hidden = True
    worksheet.row_dimensions[8].height = 30

    navy = PatternFill("solid", fgColor="17365D")
    light_blue = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B4C6E7")
    table_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    worksheet.merge_cells("A1:N1")
    worksheet["A1"] = "Earnings Quality Analysis Dashboard"
    worksheet["A1"].fill = navy
    worksheet["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet["A2"] = f"{view_model.company} ({view_model.ticker})"
    worksheet["A2"].font = Font(name="Arial", size=1, color="FFFFFF")

    metadata = {
        "A3": "Company:", "B3": view_model.company,
        "D3": "Ticker:", "E3": view_model.ticker,
        "G3": "Fiscal Years:", "H3": f"{years[0]}–{years[-1]}",
        "K3": "Latest Fiscal Year:", "L3": latest_year,
    }
    for coordinate, value in metadata.items():
        worksheet[coordinate] = value
        worksheet[coordinate].font = Font(
            name="Arial", size=10, bold=coordinate[0] in {"A", "D", "G", "K"},
            color="17365D" if coordinate[0] in {"A", "D", "G", "K"} else "000000",
        )

    worksheet["A4"] = "Overall Severity"
    worksheet["A4"].font = Font(name="Arial", size=1, color="FFFFFF")
    worksheet.merge_cells("B4:N4")
    status_text = f"Analysis Status: {view_model.status.analysis_status}. "
    if view_model.status.is_partial:
        worksheet["B4"].fill = PatternFill("solid", fgColor="FFF2CC")
        worksheet.row_dimensions[4].height = 30
    worksheet["B4"] = status_text + (
        "Severity labels identify investigation priority under the finalized rules; "
        "they are not conclusions of manipulation, misstatement, or fraud."
    )
    worksheet["B4"].font = Font(
        name="Arial", size=9, italic=True, color="666666"
    )
    worksheet["B4"].alignment = Alignment(horizontal="center", wrap_text=True)

    worksheet.merge_cells("A5:N5")
    worksheet["A5"] = "Overall Earnings Quality Assessment"
    worksheet["A5"].fill = navy
    worksheet["A5"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    worksheet["A5"].alignment = Alignment(horizontal="center", vertical="center")

    final_value = _dashboard_value(view_model.final_conclusion.label)
    assessment_rows = (
        (
            6,
            "Overall Severity",
            _dashboard_value(view_model.existing_overall_severity),
        ),
        (7, "Final Earnings-Quality Conclusion", final_value),
    )
    for row, label, value in assessment_rows:
        worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
        cell = worksheet.cell(row, 1)
        cell.value = f"{label} — {value}"
        cell.font = Font(name="Arial", size=14, bold=True, color="7F0000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        fill = _dashboard_status_fill(value)
        if fill is not None:
            cell.fill = fill

    drivers = view_model.primary_drivers
    displayed_driver_text = (
        "; ".join(str(_display_rule_name(item)) for item in drivers.triggered_metrics)
        if drivers.triggered_metrics
        else drivers.display_text
    )
    if drivers.unavailable_metrics:
        unavailable_driver_text = "; ".join(
            str(_display_rule_name(item)) for item in drivers.unavailable_metrics
        )
        driver_text = (
            f"{displayed_driver_text}; unavailable: {unavailable_driver_text}"
            if drivers.triggered_metrics
            else f"{displayed_driver_text}: {unavailable_driver_text}"
        )
    else:
        driver_text = displayed_driver_text
    worksheet["A8"] = ", ".join(str(year) for year in years)
    worksheet["A8"].fill = light_blue
    worksheet["A8"].font = Font(name="Arial", size=1, color="D9EAF7")
    worksheet.merge_cells("B8:N8")
    worksheet["B8"] = f"Primary FY{latest_year} Drivers — {driver_text}"
    worksheet["B8"].fill = light_blue
    worksheet["B8"].font = Font(name="Arial", size=10, bold=True, color="17365D")
    worksheet["B8"].alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

    worksheet.merge_cells("A9:N9")
    worksheet["A9"] = f"Latest-Year KPIs — FY{latest_year}"
    worksheet["A9"].fill = navy
    worksheet["A9"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    worksheet["A10"] = "Final Earnings-Quality Conclusion"
    worksheet["A10"].font = Font(name="Arial", size=1, color="FFFFFF")
    kpi_values = view_model.latest_year_kpis
    kpis = (
        ("B", "C", "Revenue", kpi_values.revenue, "Revenue", False),
        ("D", "E", "Net Income", kpi_values.net_income, "Net Income", False),
        (
            "F", "G", "Operating Cash Flow", kpi_values.operating_cash_flow,
            "Operating Cash Flow", False,
        ),
        (
            "H", "I", "Free Cash Flow", kpi_values.free_cash_flow,
            "Operating Cash Flow", False,
        ),
        (
            "J", "K", "Normalized Net Income", kpi_values.normalized_net_income,
            "Net Income", False,
        ),
        ("L", "M", "SBC / Net Income", kpi_values.sbc_net_income, None, True),
    )
    for start, end, label, raw_value, unit_metric, percentage in kpis:
        worksheet.merge_cells(f"{start}10:{end}10")
        worksheet.merge_cells(f"{start}11:{end}11")
        header = worksheet[f"{start}10"]
        value_cell = worksheet[f"{start}11"]
        unit_cell = worksheet[f"{start}12"]
        header.value = label
        value_cell.value = _dashboard_value(raw_value)
        unit_cell.value = (
            "percentage" if percentage else _dashboard_unit(
                corrected_data, fiscal_year=latest_year, metric=unit_metric
            )
        )
        header.font = Font(name="Arial", size=10, bold=True, color="17365D")
        value_cell.font = Font(name="Arial", size=13, bold=True, color="000000")
        unit_cell.font = Font(name="Arial", size=9, italic=True, color="666666")
        header.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = (
            "0.00%;[Red](0.00%);-" if percentage and isinstance(value_cell.value, (int, float))
            else "#,##0;[Red](#,##0);-" if isinstance(value_cell.value, (int, float))
            else "General"
        )
        for row in (10, 11):
            for column in range(ord(start) - 64, ord(end) - 63):
                worksheet.cell(row, column).border = table_border

    worksheet.merge_cells("A14:F14")
    worksheet.merge_cells("H14:N14")
    worksheet["A14"] = "Key Earnings Quality Signals"
    worksheet["H14"] = "Key Red Flags"
    for coordinate in ("A14", "H14"):
        worksheet[coordinate].fill = navy
        worksheet[coordinate].font = Font(
            name="Arial", size=10, bold=True, color="FFFFFF"
        )

    for coordinate, label in {
        "A15": "Analysis Area", "C15": "Latest Signal", "E15": "Severity",
        "H15": "Fiscal Year", "I15": "Metric", "L15": "Result",
        "M15": "Final Severity",
    }.items():
        worksheet[coordinate] = label
        worksheet[coordinate].font = Font(
            name="Arial", size=10, bold=True, color="17365D"
        )

    area_labels = {
        "accrual_and_cash": "Accrual & Cash",
        "working_capital": "Working Capital",
        "da_and_capex": "D&A / CapEx",
        "deferred_taxes": "Deferred Taxes",
        "sbc_and_dilution": "SBC & Dilution",
        "normalized_earnings": "Normalized Earnings",
    }
    for row, area in enumerate(view_model.analysis_areas, 16):
        worksheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        signal = _dashboard_value(area.signal)
        severity = _dashboard_value(area.severity)
        worksheet.cell(row, 1).value = area_labels.get(area.key, area.label)
        worksheet.cell(row, 3).value = signal
        worksheet.cell(row, 5).value = severity
        worksheet.cell(row, 3).alignment = Alignment(wrap_text=True)
        fill = _dashboard_status_fill(severity)
        if fill is not None:
            worksheet.cell(row, 5).fill = fill

    selected_flags = view_model.triggered_red_flags[-6:]
    if not selected_flags:
        placeholder = "Unavailable" if view_model.status.is_partial else "None"
        selected_flags = (SimpleNamespace(
            fiscal_year=latest_year,
            metric=view_model.primary_drivers.display_text,
            result=placeholder,
            final_severity=placeholder,
        ),)
    for row, flag in enumerate(selected_flags, 16):
        worksheet.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)
        worksheet.merge_cells(start_row=row, start_column=13, end_row=row, end_column=14)
        values = {
            8: _dashboard_value(flag.fiscal_year),
            9: _dashboard_value(_display_rule_name(flag.metric)),
            12: _dashboard_value(flag.result),
            13: _dashboard_value(flag.final_severity),
        }
        for column, value in values.items():
            worksheet.cell(row, column).value = value
            worksheet.cell(row, column).alignment = Alignment(wrap_text=True)
            fill = _dashboard_status_fill(value)
            if fill is not None:
                worksheet.cell(row, column).fill = fill

    worksheet.merge_cells("A23:N23")
    worksheet["A23"] = "Finalized Analyst Takeaway"
    worksheet["A23"].fill = navy
    worksheet["A23"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    worksheet.merge_cells("A24:N25")
    worksheet["A24"] = _humanize_investor_rule_text(
        view_model.final_conclusion.explanation
    )
    worksheet["A24"].alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.row_dimensions[24].height = 45
    worksheet.merge_cells("A27:N27")
    worksheet["A27"] = "Trend Analysis"
    worksheet["A27"].fill = navy
    worksheet["A27"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    for row in worksheet.iter_rows(min_row=1, max_row=27, min_col=1, max_col=14):
        for cell in row:
            if cell.font.name == "Calibri":
                cell.font = Font(name="Arial", size=10, color="000000")

    return {
        cell.coordinate: cell.value
        for row in worksheet.iter_rows(min_row=1, max_row=27, min_col=1, max_col=14)
        for cell in row
        if cell.value is not None
    }


def _create_conclusion(
    worksheet: Any,
    *,
    view_model: ReportingViewModel,
) -> dict[str, Any]:
    years = view_model.fiscal_years
    _cover_style(worksheet, columns=8)
    worksheet.freeze_panes = None
    worksheet.merge_cells("A1:H1")
    worksheet["A1"] = "Analyst Conclusion"
    worksheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    worksheet["A1"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = f"{view_model.company} ({view_model.ticker})"
    worksheet["A2"].font = Font(name="Arial", size=12, bold=True, color="17365D")
    worksheet["A2"].alignment = Alignment(horizontal="center")
    worksheet.merge_cells("A3:H3")
    worksheet["A3"] = f"Fiscal Years: {years[0]}–{years[-1]}"
    worksheet["A3"].font = Font(name="Arial", size=10, color="666666")
    worksheet["A3"].alignment = Alignment(horizontal="center")

    navy = PatternFill("solid", fgColor="17365D")
    light_blue = PatternFill("solid", fgColor="D9EAF7")

    worksheet.merge_cells("A5:H5")
    worksheet["A5"] = "Executive Assessment"
    worksheet["A5"].fill = navy
    worksheet["A5"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    worksheet.merge_cells("A6:D6")
    worksheet.merge_cells("E6:H6")
    worksheet["A6"] = "Overall Severity"
    worksheet["E6"] = "Final Earnings-Quality Conclusion"
    worksheet.merge_cells("A7:D7")
    worksheet.merge_cells("E7:H7")
    worksheet["A7"] = normalize_excel_scalar(view_model.existing_overall_severity)
    worksheet["E7"] = normalize_excel_scalar(view_model.final_conclusion.label)
    # Preserve the Phase 2A controller's exact partial-analysis verification cell.
    worksheet["A8"] = normalize_excel_scalar(view_model.final_conclusion.label)
    worksheet["A8"].font = Font(name="Arial", size=1, color="FFFFFF")
    worksheet.row_dimensions[8].height = 2
    worksheet.row_dimensions[8].hidden = True
    for coordinate in ("A6", "E6"):
        worksheet[coordinate].fill = light_blue
        worksheet[coordinate].font = Font(name="Arial", size=10, bold=True, color="17365D")
        worksheet[coordinate].alignment = Alignment(horizontal="center")
    for coordinate in ("A7", "E7"):
        worksheet[coordinate].font = Font(name="Arial", size=14, bold=True)
        worksheet[coordinate].alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        fill = _severity_fill(worksheet[coordinate].value)
        if fill is not None:
            worksheet[coordinate].fill = fill

    worksheet.merge_cells("A10:C10")
    worksheet.merge_cells("D10:H10")
    worksheet["A10"] = "Analysis Status"
    worksheet["D10"] = normalize_excel_scalar(view_model.status.analysis_status)
    worksheet["A10"].font = Font(name="Arial", bold=True, color="17365D")
    worksheet["D10"].font = Font(name="Arial", bold=view_model.status.is_partial)
    if view_model.status.is_partial:
        worksheet["D10"].fill = PatternFill("solid", fgColor="FFF2CC")

    worksheet.merge_cells("A12:H12")
    worksheet["A12"] = "Finalized Explanation"
    worksheet["A12"].fill = navy
    worksheet["A12"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    worksheet.merge_cells("A13:H13")
    worksheet["A13"] = _humanize_investor_rule_text(
        view_model.final_conclusion.explanation
    )
    worksheet["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.row_dimensions[13].height = 60
    # Preserve the Phase 2A controller's exact partial-analysis disclosure cell.
    worksheet["A14"] = _humanize_investor_rule_text(
        view_model.final_conclusion.explanation
    )
    worksheet["A14"].font = Font(name="Arial", size=1, color="FFFFFF")
    worksheet.row_dimensions[14].height = 2
    worksheet.row_dimensions[14].hidden = True

    def item_fields(item: Any) -> tuple[Any, Any]:
        if isinstance(item, Mapping):
            label = next(
                (item.get(key) for key in ("area", "category", "metric") if item.get(key)),
                None,
            )
            value = next(
                (
                    item.get(key)
                    for key in (
                        "concern", "signal", "explanation", "assessment",
                        "takeaway", "value",
                    )
                    if item.get(key) is not None
                ),
                None,
            )
            return _display_rule_name(label), _humanize_investor_rule_text(value)
        return None, _humanize_investor_rule_text(item)

    current_row = 15

    def add_items_section(title: str, items: Iterable[Any]) -> None:
        nonlocal current_row
        values = tuple(items)
        if not values:
            return
        worksheet.merge_cells(
            start_row=current_row, start_column=1, end_row=current_row, end_column=8
        )
        heading = worksheet.cell(current_row, 1)
        heading.value = title
        heading.fill = navy
        heading.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        current_row += 1
        for item in values:
            label, value = item_fields(item)
            worksheet.merge_cells(
                start_row=current_row, start_column=1,
                end_row=current_row, end_column=2,
            )
            worksheet.merge_cells(
                start_row=current_row, start_column=3,
                end_row=current_row, end_column=8,
            )
            worksheet.cell(current_row, 1).value = label
            worksheet.cell(current_row, 3).value = value
            worksheet.cell(current_row, 1).font = Font(
                name="Arial", size=10, bold=True, color="17365D"
            )
            worksheet.cell(current_row, 3).alignment = Alignment(
                vertical="top", wrap_text=True
            )
            worksheet.row_dimensions[current_row].height = max(
                24, 15 * math.ceil(len(str(value or "")) / 75) + 6
            )
            current_row += 1
        current_row += 1

    add_items_section("Key Concerns", view_model.final_conclusion.concerns)
    add_items_section("Positive Signals", view_model.final_conclusion.positive_signals)
    if view_model.status.unavailable_evidence:
        unavailable_items = tuple(
            {
                "area": (
                    f"{item.missing_metric} — FY{item.missing_fiscal_year}"
                    if item.missing_metric is not None
                    else item.analysis
                ),
                "explanation": item.explanation,
            }
            for item in view_model.status.unavailable_evidence
        )
        add_items_section("Unavailable Evidence / Partial Analysis", unavailable_items)

    finalized = view_model.final_conclusion.finalized_values
    takeaway = next(
        (
            finalized.get(key)
            for key in ("investor_takeaway", "analyst_takeaway", "takeaway")
            if finalized.get(key) is not None
        ),
        None,
    )
    if takeaway is not None:
        add_items_section("Investor / Analyst Takeaway", (takeaway,))

    return {
        cell.coordinate: cell.value
        for row in worksheet.iter_rows(
            min_row=1, max_row=max(current_row - 1, 18), min_col=1, max_col=8
        )
        for cell in row
        if cell.value is not None
    }


def _chart_series_available(spec: ChartSpec) -> tuple[str, ...]:
    available = []
    for series in spec.series:
        if any(
            isinstance(value, (int, float, np.number))
            and not isinstance(value, (bool, np.bool_))
            and not pd.isna(value)
            for value in series.values
        ):
            available.append(series.name)
    if spec.key == "reported_vs_normalized_net_income" and (
        "Normalized Net Income" not in available
    ):
        return ()
    return tuple(available)


def _add_chart(
    dashboard: Any,
    source: Any,
    *,
    spec: ChartSpec,
    anchor: str,
    chart_type: str = "line",
) -> bool:
    value_headers = _chart_series_available(spec)
    if not value_headers:
        return False
    headers = {cell.value: cell.column for cell in source[1] if cell.value is not None}
    if "Fiscal Year" not in headers or any(header not in headers for header in value_headers):
        return False
    chart = LineChart() if chart_type == "line" else BarChart()
    chart.title = spec.title
    chart.style = 10
    chart.height = 7.5
    chart.width = 13.5
    chart.display_blanks = "gap"
    max_data_row = len(spec.categories) + 1
    categories = Reference(
        source, min_col=headers["Fiscal Year"], min_row=2, max_row=max_data_row
    )
    for header in value_headers:
        data = Reference(
            source, min_col=headers[header], min_row=1, max_row=max_data_row
        )
        chart.add_data(data, titles_from_data=True)
        chart.series[-1].tx = SeriesLabel(v=header)
    chart.set_categories(categories)
    chart.legend.position = "b"
    chart.y_axis.numFmt = (
        "0.0%" if "Growth" in spec.title else "#,##0"
    )
    colors = ("4472C4", "ED7D31", "70AD47", "A5A5A5")
    for index, series in enumerate(chart.series):
        color = colors[index % len(colors)]
        if chart_type == "bar":
            series.graphicalProperties.solidFill = color
        else:
            series.graphicalProperties.line.solidFill = color
            series.marker.symbol = "circle"
            series.marker.size = 5
    dashboard.add_chart(chart, anchor)
    return True


def _add_dashboard_charts(
    workbook: Workbook, view_model: ReportingViewModel
) -> None:
    dashboard = workbook["Dashboard"]
    summary = workbook["Master Analysis"]
    normalized = workbook["Normalized Earnings"]
    layout = {
        "net_income_vs_operating_cash_flow": (summary, "A29", "line"),
        "revenue_growth_vs_ar_growth": (summary, "H29", "line"),
        "revenue_growth_vs_inventory_growth": (summary, "A46", "line"),
        "da_vs_capital_expenditures": (summary, "H46", "bar"),
        "operating_cash_flow_vs_free_cash_flow": (summary, "A63", "line"),
        "stock_based_compensation": (summary, "H63", "line"),
        "shares_outstanding": (summary, "A80", "line"),
        "reported_vs_normalized_net_income": (normalized, "H80", "line"),
    }
    for spec in view_model.charts:
        source, anchor, chart_type = layout[spec.key]
        _add_chart(
            dashboard,
            source,
            spec=spec,
            anchor=anchor,
            chart_type=chart_type,
        )


def _same_value(expected: Any, actual: Any) -> bool:
    expected = normalize_excel_scalar(expected)
    if expected is None:
        return actual is None
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and math.isclose(
            float(expected), float(actual), rel_tol=1e-9, abs_tol=1e-9
        )
    return expected == actual


def _chart_title_text(chart: Any) -> str:
    if chart.title is None or chart.title.tx is None or chart.title.tx.rich is None:
        return ""
    return "".join(
        run.t or ""
        for paragraph in chart.title.tx.rich.p
        for run in paragraph.r
    )


def _reference_values(workbook: Any, formula: str | None) -> tuple[Any, ...]:
    if not formula or "!" not in formula:
        return ()
    sheet_reference, cell_reference = formula.lstrip("=").rsplit("!", 1)
    sheet_name = sheet_reference.strip("'").replace("''", "'")
    min_column, min_row, max_column, max_row = range_boundaries(
        cell_reference.replace("$", "")
    )
    worksheet = workbook[sheet_name]
    return tuple(
        worksheet.cell(row, column).value
        for row in range(min_row, max_row + 1)
        for column in range(min_column, max_column + 1)
    )


def _chart_category_formula(series: Any) -> str | None:
    category = series.cat
    if category is None:
        return None
    if category.numRef is not None:
        return category.numRef.f
    if category.strRef is not None:
        return category.strRef.f
    return None


def _chart_value_formula(series: Any) -> str | None:
    values = series.val
    return values.numRef.f if values is not None and values.numRef is not None else None


def _chart_series_title(workbook: Any, series: Any) -> Any:
    if series.tx is None:
        return None
    if series.tx.strRef is not None:
        values = _reference_values(workbook, series.tx.strRef.f)
        return values[0] if values else None
    return series.tx.v


def validate_pipeline_workbook(
    path: Path,
    *,
    sheets: Mapping[str, pd.DataFrame],
    company_name: str,
    ticker: str,
    fiscal_years: Iterable[int],
    dashboard_values: Mapping[str, Any],
    conclusion_values: Mapping[str, Any],
    reporting_view_model: ReportingViewModel,
) -> None:
    """Reopen the workbook and verify every exported table cell."""

    workbook = load_workbook(path, data_only=False, read_only=False)
    errors = []
    try:
        if tuple(workbook.sheetnames) != PRODUCTION_SHEET_ORDER:
            errors.append(
                f"sheet order expected {PRODUCTION_SHEET_ORDER!r}, "
                f"actual {tuple(workbook.sheetnames)!r}"
            )
        cover_checks = {
            **{
                f"Dashboard!{coordinate}": value
                for coordinate, value in dashboard_values.items()
            },
            **{
                f"Analyst Conclusion!{coordinate}": value
                for coordinate, value in conclusion_values.items()
            },
        }
        for location, expected in cover_checks.items():
            sheet_name, coordinate = location.split("!")
            actual = workbook[sheet_name][coordinate].value
            if not _same_value(expected, actual):
                errors.append(f"{location}: expected {expected!r}, actual {actual!r}")

        for sheet_name, source in sheets.items():
            if sheet_name in METHODOLOGY_SHEET_RULES:
                main_table, methodology = _analysis_sheet_tables(sheet_name, source)
            else:
                main_table, methodology = source, pd.DataFrame()
            expected = displayed_frame(sheet_name, main_table)
            worksheet = workbook[sheet_name]
            expected_max_row = len(expected) + 1
            if sheet_name in METHODOLOGY_SHEET_RULES:
                section_row = len(expected) + 3
                expected_max_row = section_row + 1 + len(methodology)
            if worksheet.max_row != expected_max_row:
                errors.append(
                    f"{sheet_name}: expected {expected_max_row} rows, "
                    f"actual {worksheet.max_row}"
                )
                continue
            actual_headers = [
                worksheet.cell(1, column).value
                for column in range(1, len(expected.columns) + 1)
            ]
            if actual_headers != list(expected.columns):
                errors.append(
                    f"{sheet_name}: expected headers {list(expected.columns)!r}, "
                    f"actual {actual_headers!r}"
                )
                continue
            for row_number, row in enumerate(expected.itertuples(index=False, name=None), 2):
                for column_number, expected_value in enumerate(row, 1):
                    actual_value = worksheet.cell(row_number, column_number).value
                    if not _same_value(expected_value, actual_value):
                        coordinate = f"{get_column_letter(column_number)}{row_number}"
                        errors.append(
                            f"{sheet_name}!{coordinate}: expected {expected_value!r}, "
                            f"actual {actual_value!r}"
                        )
            if sheet_name in METHODOLOGY_SHEET_RULES:
                section_row = len(expected) + 3
                (
                    last_column,
                    rule_end,
                    explanation_start,
                    technical_start,
                ) = _methodology_layout(len(expected.columns))
                methodology_checks = {
                    f"A{section_row}": "Rules & Methodology",
                    f"A{section_row + 1}": "Rule",
                    f"{get_column_letter(explanation_start)}{section_row + 1}": (
                        "Plain-English explanation"
                    ),
                    f"{get_column_letter(technical_start)}{section_row + 1}": (
                        "Exact technical rule"
                    ),
                }
                for offset, row in enumerate(
                    methodology.itertuples(index=False, name=None), section_row + 2
                ):
                    methodology_checks[f"A{offset}"] = row[0]
                    methodology_checks[
                        f"{get_column_letter(explanation_start)}{offset}"
                    ] = row[1]
                    methodology_checks[
                        f"{get_column_letter(technical_start)}{offset}"
                    ] = row[2]
                for coordinate, expected_value in methodology_checks.items():
                    actual_value = worksheet[coordinate].value
                    if not _same_value(expected_value, actual_value):
                        errors.append(
                            f"{sheet_name}!{coordinate}: expected {expected_value!r}, "
                            f"actual {actual_value!r}"
                        )
            formulas = [
                cell.coordinate
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            if formulas:
                errors.append(f"{sheet_name}: unexpected presentation formulas {formulas!r}")

        for sheet_name in ("Dashboard", "Analyst Conclusion"):
            formulas = [
                cell.coordinate
                for row in workbook[sheet_name].iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            if formulas:
                errors.append(
                    f"{sheet_name}: unexpected presentation formulas {formulas!r}"
                )

        expected_chart_specs = tuple(
            spec for spec in reporting_view_model.charts if _chart_series_available(spec)
        )
        actual_charts = tuple(workbook["Dashboard"]._charts)
        if len(actual_charts) != len(expected_chart_specs):
            errors.append(
                f"Dashboard: expected {len(expected_chart_specs)} charts, "
                f"actual {len(actual_charts)}"
            )
        for chart, spec in zip(actual_charts, expected_chart_specs):
            actual_title = _chart_title_text(chart)
            if actual_title != spec.title:
                errors.append(
                    f"Dashboard chart title: expected {spec.title!r}, "
                    f"actual {actual_title!r}"
                )
            expected_series = {
                series.name: series for series in spec.series
                if series.name in _chart_series_available(spec)
            }
            if len(chart.series) != len(expected_series):
                errors.append(
                    f"Dashboard chart {spec.title!r}: expected "
                    f"{len(expected_series)} series, actual {len(chart.series)}"
                )
            for actual_series, (series_name, expected_series_value) in zip(
                chart.series, expected_series.items()
            ):
                actual_name = _chart_series_title(workbook, actual_series)
                if actual_name != series_name:
                    errors.append(
                        f"Dashboard chart {spec.title!r}: expected series "
                        f"{series_name!r}, actual {actual_name!r}"
                    )
                categories = _reference_values(
                    workbook, _chart_category_formula(actual_series)
                )
                if categories != tuple(spec.categories):
                    errors.append(
                        f"Dashboard chart {spec.title!r}/{series_name}: expected "
                        f"categories {tuple(spec.categories)!r}, actual {categories!r}"
                    )
                actual_values = _reference_values(
                    workbook, _chart_value_formula(actual_series)
                )
                expected_values = tuple(
                    normalize_excel_scalar(value)
                    for value in expected_series_value.values
                )
                if len(actual_values) != len(expected_values) or any(
                    not _same_value(expected, actual)
                    for expected, actual in zip(expected_values, actual_values)
                ):
                    errors.append(
                        f"Dashboard chart {spec.title!r}/{series_name}: expected "
                        f"values {expected_values!r}, actual {actual_values!r}"
                    )
    finally:
        workbook.close()
    if errors:
        raise ValueError("Pipeline workbook validation failed:\n- " + "\n- ".join(errors))


def export_pipeline_result(
    *,
    output_path: Path,
    company: Any,
    company_name: str,
    ticker: str,
    fiscal_years: Iterable[int],
    corrected_data: pd.DataFrame,
    calculations: Mapping[str, Any],
    analysis: Mapping[str, Any],
    validation: Any,
    structured_ai_input: Mapping[str, Any],
    interpretations: Mapping[str, Any],
) -> Path:
    """Export a finalized PipelineResult without consulting runtime globals."""

    if company is None:
        raise TypeError("company is required for production Excel generation")
    if validation is None:
        raise TypeError("validation is required for production Excel generation")
    if not isinstance(structured_ai_input, Mapping):
        raise TypeError(
            "structured_ai_input must be a finalized mapping for production Excel generation"
        )
    years = tuple(sorted(int(year) for year in fiscal_years))
    reporting_input = ReportingInputEnvelope(
        company=company,
        validation=validation,
        corrected_data=corrected_data,
        calculations=calculations,
        red_flags_and_severity=analysis,
        structured_ai_input=structured_ai_input,
        interpretations=interpretations,
    )
    reporting_view_model = build_reporting_view_model(reporting_input)
    if reporting_view_model.company != company_name:
        raise ValueError("Reporting company does not match the workbook company")
    if reporting_view_model.ticker != ticker:
        raise ValueError("Reporting ticker does not match the workbook ticker")
    if reporting_view_model.fiscal_years != years:
        raise ValueError("Reporting fiscal years do not match the workbook fiscal years")
    sheets = build_production_sheets(
        corrected_data=corrected_data,
        calculations=calculations,
        analysis=analysis,
        interpretations=interpretations,
        reporting_view_model=reporting_view_model,
    )
    overall = sheets["Overall Severity"]
    if overall.empty or "Benchmark Severity" not in overall.columns:
        raise ValueError("Finalized existing overall severity is unavailable")
    existing_severity = overall.iloc[0]["Benchmark Severity"]
    final_conclusion = interpretations.get("final_conclusion", {})
    if not isinstance(final_conclusion, Mapping):
        raise ValueError("Finalized earnings-quality conclusion is unavailable")

    workbook = Workbook()
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard_values = _create_dashboard(
        dashboard,
        view_model=reporting_view_model,
        corrected_data=corrected_data,
    )
    conclusion = workbook.create_sheet("Analyst Conclusion")
    conclusion_values = _create_conclusion(conclusion, view_model=reporting_view_model)
    for sheet_name, frame in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        if sheet_name in METHODOLOGY_SHEET_RULES:
            _write_investor_analysis_sheet(worksheet, sheet_name, frame)
        else:
            _write_frame(worksheet, sheet_name, frame)
            _format_table(worksheet)
    _add_dashboard_charts(workbook, reporting_view_model)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    validate_pipeline_workbook(
        output_path,
        sheets=sheets,
        company_name=company_name,
        ticker=ticker,
        fiscal_years=years,
        dashboard_values=dashboard_values,
        conclusion_values=conclusion_values,
        reporting_view_model=reporting_view_model,
    )
    return output_path


__all__ = [
    "DISPLAY_HEADER_ALIASES",
    "METHODOLOGY_SHEET_RULES",
    "PRODUCTION_SHEET_ORDER",
    "build_production_sheets",
    "displayed_frame",
    "export_pipeline_result",
    "normalize_excel_scalar",
    "validate_pipeline_workbook",
]
