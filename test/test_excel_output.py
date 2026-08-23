"""Integration coverage for the finalized PipelineResult Excel presentation."""

from __future__ import annotations

import math
import sys
from copy import deepcopy
from datetime import date
from inspect import Parameter, signature
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import main  # noqa: E402
import Analysis.reporting_view_model as reporting_view_model  # noqa: E402
from Analysis.excel_export import (  # noqa: E402
    DISPLAY_HEADER_ALIASES,
    METHODOLOGY_SHEET_RULES,
    PRODUCTION_SHEET_ORDER,
    build_production_sheets,
    displayed_frame,
    export_pipeline_result,
    normalize_excel_scalar,
    _analysis_sheet_tables,
    _chart_series_available,
    _INVESTOR_RULE_TEXT_ALIASES,
    _methodology_layout,
)
from Analysis.reporting_view_model import RULE_PRESENTATION_MAPPING  # noqa: E402
from Analysis.manual_one_off_classification import (  # noqa: E402
    manual_one_off_classification_df,
)
from Data.analysis_validation_gate import (  # noqa: E402
    AnalysisDecision,
    AnalysisValidationGateResult,
    BlockingIssue,
)
from Data.apple_test_dataset import apple_financial_data_df  # noqa: E402
from Data.financial_statement_fetcher import FactStatus  # noqa: E402


APPLE = SimpleNamespace(company_name="Apple Inc.", ticker="AAPL")
DISNEY = SimpleNamespace(company_name="The Walt Disney Company", ticker="DIS")
OTHER = SimpleNamespace(company_name="Example Industries", ticker="EXM")


def _five_year_apple() -> pd.DataFrame:
    prior = apple_financial_data_df.loc[
        apple_financial_data_df["Fiscal Year"].eq(2022)
    ].copy(deep=True)
    prior["Fiscal Year"] = 2021
    return pd.concat(
        [prior, apple_financial_data_df.copy(deep=True)], ignore_index=True
    )


def _five_year_one_off_fixture() -> pd.DataFrame:
    prior = manual_one_off_classification_df.loc[
        manual_one_off_classification_df["Fiscal Year"].eq(2022)
    ].copy(deep=True)
    prior["Fiscal Year"] = 2021
    prior["Category"] = "Unavailable"
    prior["Description"] = "Unavailable"
    prior["Value"] = float("nan")
    prior["Units"] = "Not Available"
    prior["Direction"] = "Not Available"
    prior["Tax Basis"] = "Not Available"
    prior["Availability"] = "Not Available"
    prior["Source"] = "Not Available"
    prior["Source Date"] = "Not Available"
    prior["Source Detail"] = "Not Available"
    prior["Manual Classification"] = "Not Applicable"
    prior["Manual Review Status"] = "Not Applicable"
    return pd.concat(
        [prior, manual_one_off_classification_df.copy(deep=True)], ignore_index=True
    )


def _company_data(frame: pd.DataFrame, company) -> pd.DataFrame:
    result = frame.copy(deep=True)
    result["Company"] = company.company_name
    result["Ticker"] = company.ticker
    if "Source" in result.columns:
        result["Source"] = "SEC Form 10-K"
    return result


def _five_year_disney() -> pd.DataFrame:
    result = _company_data(_five_year_apple(), DISNEY)
    latest_values = {
        "Revenue": 125000.0,
        "Net Income": 9500.0,
        "Operating Cash Flow": 14000.0,
        "Capital Expenditures": 6000.0,
        "Stock-Based Compensation": 1800.0,
    }
    for metric, value in latest_values.items():
        result.loc[
            result["Fiscal Year"].eq(2025) & result["Metric"].eq(metric),
            "Value",
        ] = value
    return result


def _gate(frame: pd.DataFrame) -> AnalysisValidationGateResult:
    years = tuple(sorted(int(year) for year in frame["Fiscal Year"].unique()))
    return AnalysisValidationGateResult(
        decision=AnalysisDecision.CONTINUE,
        can_analyze=True,
        requested_fiscal_years=years,
        blocking_issues=(),
        metric_year_results=(),
        blocking_issue_count=0,
    )


def _partial_gate(frame: pd.DataFrame) -> AnalysisValidationGateResult:
    years = tuple(sorted(int(year) for year in frame["Fiscal Year"].unique()))
    issue = BlockingIssue(
        metric="Inventory",
        fiscal_year=years[-1],
        validation_source="Focused same-run reporting transport test",
        blocking_status="MISSING",
        explanation="Required value is unavailable.",
        manual_correction_used=False,
    )
    return AnalysisValidationGateResult(
        decision=AnalysisDecision.CONTINUE,
        can_analyze=True,
        requested_fiscal_years=years,
        blocking_issues=(issue,),
        metric_year_results=(),
        blocking_issue_count=1,
    )


def _finalized_run(
    monkeypatch,
    tmp_path,
    frame: pd.DataFrame,
    company,
    *,
    one_off_evidence=None,
    validation=None,
):
    monkeypatch.setattr(main, "_PROJECT_ROOT", PROJECT_ROOT)
    modules = main.create_production_pipeline_modules(
        user_agent="EarningsQualityAgent tests@example.com",
        one_off_evidence=(
            None if one_off_evidence is None else one_off_evidence.copy(deep=True)
        ),
    )
    corrected = frame.copy(deep=True)
    validation = validation or _gate(corrected)
    calculations = modules.run_calculations(corrected)
    analysis = modules.run_red_flags_and_severity(calculations)
    structured = modules.build_structured_ai_input(
        company, validation, corrected, calculations, analysis
    )
    interpretations = modules.run_interpretations(structured)
    monkeypatch.setattr(main, "_PROJECT_ROOT", tmp_path)
    path = main._production_excel_output(
        company,
        corrected,
        calculations,
        analysis,
        structured,
        interpretations,
        validation,
    )
    return SimpleNamespace(
        company=company,
        validation=validation,
        corrected_data=corrected,
        calculations=calculations,
        analysis=analysis,
        structured=structured,
        interpretations=interpretations,
        path=path,
    )


def _same_value(expected, actual) -> bool:
    if expected is None:
        return actual is None
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and math.isclose(
            float(expected), float(actual), rel_tol=1e-9, abs_tol=1e-9
        )
    return expected == actual


def _reporting_view(run):
    return reporting_view_model.build_reporting_view_model(
        reporting_view_model.ReportingInputEnvelope(
            company=run.company,
            validation=run.validation,
            corrected_data=run.corrected_data,
            calculations=run.calculations,
            red_flags_and_severity=run.analysis,
            structured_ai_input=run.structured,
            interpretations=run.interpretations,
        )
    )


def _assert_sheet_matches_frame(workbook, sheet_name, source):
    if sheet_name in METHODOLOGY_SHEET_RULES:
        main_table, methodology = _analysis_sheet_tables(sheet_name, source)
    else:
        main_table, methodology = source, pd.DataFrame()
    expected = displayed_frame(sheet_name, main_table)
    worksheet = workbook[sheet_name]
    expected_max_row = len(expected) + 1
    if not methodology.empty:
        expected_max_row = len(expected) + 4 + len(methodology)
    assert worksheet.max_row == expected_max_row
    assert [
        worksheet.cell(1, column).value
        for column in range(1, len(expected.columns) + 1)
    ] == list(expected.columns)
    for row_number, row in enumerate(expected.itertuples(index=False, name=None), 2):
        for column_number, expected_value in enumerate(row, 1):
            assert _same_value(
                expected_value, worksheet.cell(row_number, column_number).value
            ), f"{sheet_name}!{worksheet.cell(row_number, column_number).coordinate}"
    if not methodology.empty:
        section_row = len(expected) + 3
        _, _, explanation_start, technical_start = _methodology_layout(
            len(expected.columns)
        )
        assert worksheet[f"A{section_row}"].value == "Rules & Methodology"
        assert worksheet[f"A{section_row + 1}"].value == "Rule"
        assert worksheet.cell(
            section_row + 1, explanation_start
        ).value == "Plain-English explanation"
        assert worksheet.cell(
            section_row + 1, technical_start
        ).value == "Exact technical rule"
        assert [
            (
                worksheet.cell(row, 1).value,
                worksheet.cell(row, explanation_start).value,
                worksheet.cell(row, technical_start).value,
            )
            for row in range(section_row + 2, worksheet.max_row + 1)
        ] == list(methodology.itertuples(index=False, name=None))


def _chart_title(chart):
    return "".join(
        run.t or ""
        for paragraph in chart.title.tx.rich.p
        for run in paragraph.r
    )


def _chart_reference_values(workbook, formula):
    sheet_reference, cell_reference = formula.rsplit("!", 1)
    sheet_name = sheet_reference.strip("'").replace("''", "'")
    worksheet = workbook[sheet_name]
    cells = worksheet[cell_reference.replace("$", "")]
    if not isinstance(cells, tuple):
        return (cells.value,)
    return tuple(cell.value for row in cells for cell in (row if isinstance(row, tuple) else (row,)))


def _workbook_snapshot(path: Path):
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        return {
            sheet_name: tuple(
                tuple(cell.value for cell in row)
                for row in workbook[sheet_name].iter_rows()
            )
            for sheet_name in workbook.sheetnames
        }, len(workbook["Dashboard"]._charts)
    finally:
        workbook.close()


def _assert_approved_methodology_sections(workbook, sheets) -> None:
    for sheet_name, rule_keys in METHODOLOGY_SHEET_RULES.items():
        source = sheets[sheet_name]
        main_table, methodology = _analysis_sheet_tables(sheet_name, source)
        worksheet = workbook[sheet_name]
        section_row = len(main_table) + 3
        header_row = section_row + 1
        _, rule_end, explanation_start, technical_start = _methodology_layout(
            len(main_table.columns)
        )

        assert all(
            worksheet.cell(section_row - 1, column).value is None
            for column in range(1, worksheet.max_column + 1)
        )
        assert worksheet.cell(section_row - 2, 1).value is not None
        section = worksheet.cell(section_row, 1)
        assert section.value == "Rules & Methodology"
        assert section.fill.fgColor.rgb == "0017365D"
        assert section.font.color.rgb == "00FFFFFF"
        assert (worksheet.cell(header_row, 1).value,) == ("Rule",)
        assert worksheet.cell(
            header_row, explanation_start
        ).value == "Plain-English explanation"
        assert worksheet.cell(
            header_row, technical_start
        ).value == "Exact technical rule"
        assert (technical_start - explanation_start) > rule_end

        expected_rows = [
            (
                RULE_PRESENTATION_MAPPING[key].display_name,
                RULE_PRESENTATION_MAPPING[key].plain_english_explanation,
                RULE_PRESENTATION_MAPPING[key].exact_technical_rule,
            )
            for key in rule_keys
        ]
        actual_rows = []
        for row in range(header_row + 1, worksheet.max_row + 1):
            actual_rows.append(
                (
                    worksheet.cell(row, 1).value,
                    worksheet.cell(row, explanation_start).value,
                    worksheet.cell(row, technical_start).value,
                )
            )
            assert worksheet.cell(row, explanation_start).alignment.wrap_text
            assert worksheet.cell(row, technical_start).alignment.wrap_text
            assert worksheet.row_dimensions[row].height >= 42
        assert actual_rows == expected_rows
        assert len(actual_rows) == len(set(actual_rows)) == len(rule_keys)

    red_flags = workbook["Red Flags"]
    headers = {cell.value: cell.column for cell in red_flags[1]}
    internal_metrics = set(RULE_PRESENTATION_MAPPING)
    for row in range(2, 2 + len(sheets["Red Flags"])):
        displayed_metric = red_flags.cell(row, headers["Metric"]).value
        explanation = red_flags.cell(row, headers["Rule"]).value
        source_metric = sheets["Red Flags"].iloc[row - 2]["Metric"]
        assert displayed_metric == RULE_PRESENTATION_MAPPING[source_metric].display_name
        assert explanation == RULE_PRESENTATION_MAPPING[
            source_metric
        ].plain_english_explanation
        assert displayed_metric not in internal_metrics or (
            displayed_metric
            == RULE_PRESENTATION_MAPPING[source_metric].display_name
            == source_metric
        )


def test_excel_scalar_normalization_and_strict_reopen_validation(
    monkeypatch, tmp_path
):
    sample_date = date(2025, 9, 27)
    assert normalize_excel_scalar(FactStatus.RETRIEVED) == "RETRIEVED"
    assert normalize_excel_scalar(None) is None
    assert normalize_excel_scalar("Unavailable") == "Unavailable"
    assert normalize_excel_scalar("existing text") == "existing text"
    assert normalize_excel_scalar(42) == 42
    assert normalize_excel_scalar(sample_date) == sample_date
    assert normalize_excel_scalar(True) == "Yes"
    assert normalize_excel_scalar(False) == "No"

    frame = _five_year_apple()
    frame["Status"] = FactStatus.RETRIEVED
    frame["Optional Evidence"] = None
    frame.loc[frame.index[0], "Optional Evidence"] = "Unavailable"

    run = _finalized_run(
        monkeypatch,
        tmp_path,
        frame,
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )

    workbook = load_workbook(run.path, data_only=False, read_only=False)
    try:
        worksheet = workbook["Raw Financial Data"]
        headers = {
            worksheet.cell(1, column).value: column
            for column in range(1, worksheet.max_column + 1)
        }
        status_values = {
            worksheet.cell(row, headers["Status"]).value
            for row in range(2, worksheet.max_row + 1)
        }
        assert status_values == {FactStatus.RETRIEVED.value}
        assert (
            worksheet.cell(2, headers["Optional Evidence"]).value
            == "Unavailable"
        )
        assert all(
            worksheet.cell(row, headers["Optional Evidence"]).value is None
            for row in range(3, worksheet.max_row + 1)
        )
    finally:
        workbook.close()


def test_exact_same_run_reporting_contract_builds_view_model_without_mutation(
    monkeypatch, tmp_path
):
    captured = {}
    real_builder = reporting_view_model.build_reporting_view_model

    def capture_builder(envelope):
        before = {
            "corrected_data": envelope.corrected_data.copy(deep=True),
            "financial_summary": envelope.calculations["financial_summary"].copy(
                deep=True
            ),
            "red_flags": envelope.red_flags_and_severity["red_flags"].copy(
                deep=True
            ),
            "cash_conversion": envelope.calculations["cash_conversion"].copy(
                deep=True
            ),
            "working_capital": envelope.red_flags_and_severity[
                "working_capital"
            ].copy(deep=True),
            "capex": envelope.calculations["excel_support_tables"]["capex"].copy(
                deep=True
            ),
            "taxes": envelope.red_flags_and_severity["taxes"].copy(deep=True),
            "sbc": envelope.red_flags_and_severity["sbc"].copy(deep=True),
            "normalized_earnings": envelope.calculations[
                "normalized_earnings"
            ].copy(deep=True),
            "structured_ai_input": deepcopy(envelope.structured_ai_input),
            "interpretations": deepcopy(envelope.interpretations),
        }
        view_model = real_builder(envelope)
        captured.update(
            envelope=envelope,
            view_model=view_model,
            before=before,
        )
        return view_model

    monkeypatch.setattr(
        reporting_view_model, "build_reporting_view_model", capture_builder
    )
    frame = _company_data(_five_year_apple(), OTHER)
    frame.loc[
        frame["Metric"].eq("Inventory") & frame["Fiscal Year"].eq(2025),
        "Value",
    ] = None
    validation = _partial_gate(frame)
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        frame,
        OTHER,
        validation=validation,
    )

    envelope = captured["envelope"]
    view_model = captured["view_model"]
    assert envelope.company is run.company
    assert envelope.validation is run.validation is validation
    assert envelope.corrected_data is run.corrected_data
    assert envelope.calculations is run.calculations
    assert envelope.red_flags_and_severity is run.analysis
    assert envelope.structured_ai_input is run.structured
    assert envelope.interpretations is run.interpretations
    assert view_model.company == OTHER.company_name
    assert view_model.ticker == OTHER.ticker
    assert view_model.fiscal_years == validation.requested_fiscal_years
    assert view_model.status.validation_decision == AnalysisDecision.CONTINUE.value
    assert view_model.status.can_analyze is True
    assert view_model.status.is_partial is True
    assert view_model.status.analysis_status == "Partial Analysis"

    workbook = load_workbook(run.path, data_only=False, read_only=True)
    try:
        assert "Analysis Status: Partial Analysis" in workbook["Dashboard"]["B4"].value
        assert workbook["Analyst Conclusion"]["D10"].value == "Partial Analysis"
        assert any(
            cell.value == "Unavailable Evidence / Partial Analysis"
            for row in workbook["Analyst Conclusion"].iter_rows()
            for cell in row
        )
    finally:
        workbook.close()

    before = captured["before"]
    pd.testing.assert_frame_equal(run.corrected_data, before["corrected_data"])
    pd.testing.assert_frame_equal(
        run.calculations["financial_summary"], before["financial_summary"]
    )
    pd.testing.assert_frame_equal(run.analysis["red_flags"], before["red_flags"])
    pd.testing.assert_frame_equal(
        run.calculations["cash_conversion"], before["cash_conversion"]
    )
    pd.testing.assert_frame_equal(
        run.analysis["working_capital"], before["working_capital"]
    )
    pd.testing.assert_frame_equal(
        run.calculations["excel_support_tables"]["capex"], before["capex"]
    )
    pd.testing.assert_frame_equal(run.analysis["taxes"], before["taxes"])
    pd.testing.assert_frame_equal(run.analysis["sbc"], before["sbc"])
    pd.testing.assert_frame_equal(
        run.calculations["normalized_earnings"], before["normalized_earnings"]
    )
    assert run.structured == before["structured_ai_input"]
    assert run.interpretations == before["interpretations"]


def test_production_export_requires_complete_reporting_contract(tmp_path):
    required = {
        "output_path": tmp_path / "not-created.xlsx",
        "company": APPLE,
        "company_name": APPLE.company_name,
        "ticker": APPLE.ticker,
        "fiscal_years": (2022, 2023, 2024, 2025),
        "corrected_data": pd.DataFrame(),
        "calculations": {},
        "analysis": {},
        "validation": object(),
        "structured_ai_input": {},
        "interpretations": {},
    }
    parameters = signature(export_pipeline_result).parameters
    for argument in ("company", "validation", "structured_ai_input"):
        assert parameters[argument].default is Parameter.empty
        incomplete = {key: value for key, value in required.items() if key != argument}
        with pytest.raises(TypeError, match=argument):
            export_pipeline_result(**incomplete)


def test_every_exported_calculation_and_analysis_field_matches_pipeline_result(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    expected_sheets = build_production_sheets(
        corrected_data=run.corrected_data,
        calculations=run.calculations,
        analysis=run.analysis,
        interpretations=run.interpretations,
        reporting_view_model=_reporting_view(run),
    )
    workbook = load_workbook(run.path, data_only=False, read_only=False)
    try:
        assert tuple(workbook.sheetnames) == PRODUCTION_SHEET_ORDER
        for sheet_name, source in expected_sheets.items():
            _assert_sheet_matches_frame(workbook, sheet_name, source)
        _assert_approved_methodology_sections(workbook, expected_sheets)
        analyst_conclusion_text = " ".join(
            str(cell.value)
            for row in workbook["Analyst Conclusion"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert all(
            obsolete_text not in analyst_conclusion_text
            for obsolete_text, _ in _INVESTOR_RULE_TEXT_ALIASES
        )
        assert "AR-versus-revenue" in str(run.interpretations["final_conclusion"])
        assert not any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for sheet_name in expected_sheets
            for row in workbook[sheet_name].iter_rows()
            for cell in row
        )
        assert len(workbook["Dashboard"]._charts) == 8
    finally:
        workbook.close()


def test_investor_analysis_sheets_match_same_run_reporting_view_model(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    view_model = _reporting_view(run)
    expected_sheets = build_production_sheets(
        corrected_data=run.corrected_data,
        calculations=run.calculations,
        analysis=run.analysis,
        interpretations=run.interpretations,
        reporting_view_model=view_model,
    )
    specs = {
        "Accrual & Cash Analysis": (
            "accrual_and_cash",
            ["Fiscal Year", "Net Income Growth", "OCF Growth", "Revenue Growth"],
        ),
        "Working Capital": (
            "working_capital",
            [
                "Fiscal Year", "Inventory Growth", "Revenue Growth",
                "Inventory Revenue Gap",
            ],
        ),
        "D&A and CapEx": (
            "da_and_capex",
            [
                "Fiscal Year", "CapEx", "D&A", "CapEx / D&A",
            ],
        ),
        "Deferred Taxes": (
            "deferred_taxes",
            [
                "Fiscal Year", "Deferred Tax Assets", "Deferred Tax Liabilities",
                "Net Income",
            ],
        ),
        "SBC & Dilution": (
            "sbc_and_dilution",
            [
                "Fiscal Year", "SBC", "Net Income", "SBC / Net Income",
            ],
        ),
        "Normalized Earnings": (
            "normalized_earnings",
            [
                "Fiscal Year", "Reported Net Income", "Normalized Net Income",
                "Difference",
            ],
        ),
    }

    workbook = load_workbook(run.path, data_only=False, read_only=True)
    try:
        for sheet_name, (area_key, leading_headers) in specs.items():
            worksheet = workbook[sheet_name]
            headers = [cell.value for cell in worksheet[1]]
            assert headers[:4] == leading_headers
            header_columns = {
                value: index for index, value in enumerate(headers, start=1)
            }
            area = view_model.area(area_key)
            rows_by_year = {
                worksheet.cell(row, header_columns["Fiscal Year"]).value: row
                for row in range(2, len(area.finalized_rows) + 2)
            }
            assessments = {
                item.fiscal_year: item for item in area.yearly_assessments
            }
            aliases = DISPLAY_HEADER_ALIASES.get(sheet_name, {})
            for finalized_row in area.finalized_rows:
                fiscal_year = finalized_row["Fiscal Year"]
                workbook_row = rows_by_year[fiscal_year]
                for source_column in area.source_columns:
                    if "rule" in source_column.casefold():
                        continue
                    displayed_header = aliases.get(source_column, source_column)
                    assert _same_value(
                        normalize_excel_scalar(finalized_row[source_column]),
                        worksheet.cell(
                            workbook_row, header_columns[displayed_header]
                        ).value,
                    ), f"{sheet_name} FY{fiscal_year} {source_column}"
                expected_interpretation = (
                    assessments[fiscal_year].assessment
                    if fiscal_year in assessments
                    and assessments[fiscal_year].assessment is not None
                    else None
                )
                assert worksheet.cell(
                    workbook_row, header_columns["Analyst Interpretation"]
                ).value == expected_interpretation
            main_table, methodology = _analysis_sheet_tables(
                sheet_name, expected_sheets[sheet_name]
            )
            _, _, explanation_start, technical_start = _methodology_layout(
                len(main_table.columns)
            )
            methodology_rows = [
                (
                    worksheet.cell(row, 1).value,
                    worksheet.cell(row, explanation_start).value,
                    worksheet.cell(row, technical_start).value,
                )
                for row in range(
                    len(area.finalized_rows) + 5, worksheet.max_row + 1
                )
            ]
            assert methodology_rows == list(
                methodology.itertuples(index=False, name=None)
            )
            assert len(methodology_rows) == len(set(methodology_rows))
            if area_key == "da_and_capex":
                for fiscal_year, assessment in assessments.items():
                    assert worksheet.cell(
                        rows_by_year[fiscal_year], header_columns["Severity"]
                    ).value == normalize_excel_scalar(assessment.existing_severity)

        unrelated_sources = {
            "Raw Financial Data": run.corrected_data,
            "Master Analysis": run.calculations["financial_summary"],
            "Analyst Summary": run.analysis["fiscal_year_severity"],
            "Detailed Severity": run.analysis["detailed_severity"],
            "Overall Severity": run.analysis["overall_severity"],
            "Unavailable Outputs": run.analysis["unavailable_outputs"],
        }
        for sheet_name, source in unrelated_sources.items():
            _assert_sheet_matches_frame(workbook, sheet_name, source)
    finally:
        workbook.close()


def test_dashboard_restores_original_sections_layout_and_compact_styles(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    workbook = load_workbook(run.path, data_only=False, read_only=False)
    try:
        dashboard = workbook["Dashboard"]
        assert dashboard["A1"].value == "Earnings Quality Analysis Dashboard"
        assert (dashboard["A3"].value, dashboard["D3"].value) == (
            "Company:", "Ticker:"
        )
        assert (dashboard["G3"].value, dashboard["K3"].value) == (
            "Fiscal Years:", "Latest Fiscal Year:"
        )
        assert "investigation priority" in dashboard["B4"].value
        assert dashboard["A5"].value == "Overall Earnings Quality Assessment"
        assert dashboard["A6"].value.startswith("Overall Severity — ")
        assert dashboard["A7"].value.startswith(
            "Final Earnings-Quality Conclusion — "
        )
        assert dashboard["B8"].value.startswith("Primary FY2025 Drivers — ")
        assert dashboard["A9"].value == "Latest-Year KPIs — FY2025"
        assert dashboard["A14"].value == "Key Earnings Quality Signals"
        assert dashboard["H14"].value == "Key Red Flags"

        merges = {str(merged) for merged in dashboard.merged_cells.ranges}
        expected_merges = {
            "A1:N1", "B4:N4", "A5:N5", "A6:N6", "A7:N7", "B8:N8",
            "A9:N9", "B10:C10", "B11:C11", "L10:M10", "L11:M11",
            "A14:F14", "H14:N14", "C16:D16", "I16:K16", "M16:N16",
            "C21:D21", "I21:K21", "M21:N21",
        }
        assert expected_merges.issubset(merges)
        for coordinate in ("A1", "A5", "A9", "A14", "H14"):
            assert dashboard[coordinate].fill.fgColor.rgb == "0017365D"
            assert dashboard[coordinate].font.bold
        assert dashboard["A1"].font.name == "Arial"
        assert dashboard["A1"].font.sz == 18
        assert dashboard["A1"].font.color.rgb == "00FFFFFF"
        assert dashboard["A6"].fill.fgColor.rgb == "00F4CCCC"
        assert dashboard.sheet_view.showGridLines is False
        assert dashboard.freeze_panes is None
        assert dashboard.row_dimensions[1].height == 30
        assert dashboard.row_dimensions[8].height == 30
        assert all(
            dashboard.column_dimensions[column].width == 14
            for column in "ABCDEFGHIJKLMN"
        )
        assert dashboard.max_row == 27
        assert dashboard.max_column == 14
        blank_rows = [
            row
            for row in range(1, 28)
            if not any(
                dashboard.cell(row, column).value is not None
                for column in range(1, 15)
            )
        ]
        assert blank_rows == [13, 22, 25, 26]
        assert len(dashboard._charts) == 8
        assert dashboard["A23"].value == "Finalized Analyst Takeaway"
        assert dashboard["A27"].value == "Trend Analysis"
    finally:
        workbook.close()


def test_dashboard_charts_use_exact_selected_years_and_finalized_series(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    view_model = _reporting_view(run)
    expected_specs = [
        spec for spec in view_model.charts if _chart_series_available(spec)
    ]
    workbook = load_workbook(run.path, data_only=False, read_only=False)
    try:
        charts = workbook["Dashboard"]._charts
        assert [_chart_title(chart) for chart in charts] == [
            spec.title for spec in expected_specs
        ]
        for chart, spec in zip(charts, expected_specs):
            expected_series = [
                series
                for series in spec.series
                if series.name in _chart_series_available(spec)
            ]
            assert len(chart.series) == len(expected_series)
            for actual, expected in zip(chart.series, expected_series):
                assert actual.tx.v == expected.name
                category_formula = (
                    actual.cat.numRef.f
                    if actual.cat.numRef is not None
                    else actual.cat.strRef.f
                )
                assert _chart_reference_values(
                    workbook, category_formula
                ) == tuple(view_model.fiscal_years)
                actual_values = _chart_reference_values(
                    workbook, actual.val.numRef.f
                )
                expected_values = tuple(
                    normalize_excel_scalar(value) for value in expected.values
                )
                assert len(actual_values) == len(expected_values)
                assert all(
                    _same_value(expected_value, actual_value)
                    for expected_value, actual_value in zip(
                        expected_values, actual_values
                    )
                )
    finally:
        workbook.close()


def test_unavailable_normalization_exports_no_apple_benchmark_content(
    monkeypatch, tmp_path
):
    other_data = _company_data(apple_financial_data_df, OTHER)
    run = _finalized_run(monkeypatch, tmp_path, other_data, OTHER)
    view_model = _reporting_view(run)
    workbook = load_workbook(run.path, data_only=False, read_only=False)
    try:
        text = " ".join(
            str(cell.value)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        ).casefold()
        for apple_content in (
            "apple", "aapl", "state aid", "10200", "10,200",
            "unusual tax gains/losses",
        ):
            assert apple_content not in text
        assert not any(
            row[0] == "normalized_earnings"
            for row in workbook["Analytical Explanations"].iter_rows(
                min_row=2, values_only=True
            )
        )
        normalized = workbook["Normalized Earnings"]
        headers = {cell.value: cell.column for cell in normalized[1]}
        assert all(
            normalized.cell(row, headers["Normalization Signal"]).value
            == "Unavailable"
            for row in range(2, len(view_model.fiscal_years) + 2)
        )
        assessments = {
            item.fiscal_year: item
            for item in view_model.area("normalized_earnings").yearly_assessments
        }
        assert all(
            normalized.cell(row, headers["Analyst Interpretation"]).value
            == assessments[
                normalized.cell(row, headers["Fiscal Year"]).value
            ].assessment
            for row in range(2, len(view_model.fiscal_years) + 2)
        )
        assert workbook["Dashboard"]["J11"].value == "Unavailable"
        assert workbook["Dashboard"]["C21"].value == "Unavailable"
        assert workbook["Dashboard"]["E21"].value == "Unavailable"
        red_flags = workbook["Red Flags"]
        red_flag_headers = {cell.value: cell.column for cell in red_flags[1]}
        repeated_rows = [
            row
            for row in range(2, red_flags.max_row + 1)
            if red_flags.cell(
                row, red_flag_headers["Metric"]
            ).value == "Repeated unusual items"
        ]
        assert len(repeated_rows) == 3
        for row in repeated_rows:
            for field in ("Raw Data", "Calculation", "Result", "Severity"):
                assert red_flags.cell(
                    row, red_flag_headers[field]
                ).value == "Unavailable"
            assert "not independently available" in red_flags.cell(
                row, red_flag_headers["Explanation"]
            ).value
        chart_titles = [_chart_title(chart) for chart in workbook["Dashboard"]._charts]
        assert "Reported vs Normalized Net Income" not in chart_titles
        assert len(chart_titles) == 7
        assert all(
            value is None
            for spec in view_model.charts
            if spec.key == "reported_vs_normalized_net_income"
            for series in spec.series
            if series.name == "Normalized Net Income"
            for value in series.values
        )
    finally:
        workbook.close()


def test_apple_disney_apple_dashboard_values_are_own_run_and_deterministic(
    monkeypatch, tmp_path
):
    reporting_runs = []
    real_builder = reporting_view_model.build_reporting_view_model

    def capture_builder(envelope):
        view_model = real_builder(envelope)
        reporting_runs.append((envelope, view_model))
        return view_model

    monkeypatch.setattr(
        reporting_view_model, "build_reporting_view_model", capture_builder
    )
    first = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    first_snapshot = _workbook_snapshot(first.path)
    apple_latest = first.calculations["financial_summary"].iloc[-1]
    first_workbook = load_workbook(first.path, data_only=False, read_only=True)
    try:
        assert first_workbook["Dashboard"]["B3"].value == APPLE.company_name
        assert first_workbook["Dashboard"]["E3"].value == APPLE.ticker
        assert first_workbook["Dashboard"]["B11"].value == apple_latest["Revenue"]
        assert first_workbook["Dashboard"]["C16"].value == apple_latest[
            "Accrual Signal"
        ]
    finally:
        first_workbook.close()

    disney = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_disney(),
        DISNEY,
    )
    disney_latest = disney.calculations["financial_summary"].iloc[-1]
    disney_workbook = load_workbook(disney.path, data_only=False, read_only=True)
    try:
        assert disney_workbook["Dashboard"]["B3"].value == DISNEY.company_name
        assert disney_workbook["Dashboard"]["E3"].value == DISNEY.ticker
        assert disney_workbook["Dashboard"]["B11"].value == disney_latest["Revenue"]
        assert disney_workbook["Dashboard"]["D11"].value == disney_latest[
            "Net Income"
        ]
        assert disney_workbook["Dashboard"]["C16"].value == disney_latest[
            "Accrual Signal"
        ]
        assert disney_workbook["Dashboard"]["J11"].value == "Unavailable"
        accrual_sheet = disney_workbook["Accrual & Cash Analysis"]
        accrual_headers = {cell.value: cell.column for cell in accrual_sheet[1]}
        disney_assessment = reporting_runs[1][1].area(
            "accrual_and_cash"
        ).yearly_assessments[-1]
        latest_row = next(
            row
            for row in range(
                2,
                len(reporting_runs[1][1].area("accrual_and_cash").finalized_rows) + 2,
            )
            if accrual_sheet.cell(row, accrual_headers["Fiscal Year"]).value == 2025
        )
        assert accrual_sheet.cell(
            latest_row,
            accrual_headers["Analyst Interpretation"],
        ).value == disney_assessment.assessment
    finally:
        disney_workbook.close()

    second = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )

    assert _workbook_snapshot(second.path) == first_snapshot
    assert [view.company for _, view in reporting_runs] == [
        APPLE.company_name,
        DISNEY.company_name,
        APPLE.company_name,
    ]
    assert [view.ticker for _, view in reporting_runs] == ["AAPL", "DIS", "AAPL"]
    assert reporting_runs[0][0].validation is first.validation
    assert reporting_runs[1][0].validation is disney.validation
    assert reporting_runs[2][0].validation is second.validation
    assert reporting_runs[0][1] is not reporting_runs[2][1]
    assert reporting_runs[0][0] is not reporting_runs[2][0]
    assert reporting_runs[0][1].area("accrual_and_cash").finalized_rows == (
        reporting_runs[2][1].area("accrual_and_cash").finalized_rows
    )
    assert reporting_runs[0][1].area("accrual_and_cash").finalized_rows != (
        reporting_runs[1][1].area("accrual_and_cash").finalized_rows
    )


def test_all_exported_analytical_explanations_are_finalized_explanations(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    workbook = load_workbook(run.path, data_only=False, read_only=True)
    try:
        rows = {
            (area, field): value
            for area, field, value in workbook["Analytical Explanations"].iter_rows(
                min_row=2, values_only=True
            )
        }
        for area, interpretation in run.interpretations.items():
            assert rows[(area, "explanation")] == interpretation["explanation"]
    finally:
        workbook.close()


def test_severity_and_final_conclusion_are_distinct_and_correct(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    severity = run.analysis["overall_severity"].iloc[0]["Benchmark Severity"]
    conclusion = run.interpretations["final_conclusion"][
        "earnings_quality_conclusion"
    ]
    workbook = load_workbook(run.path, data_only=False, read_only=True)
    try:
        dashboard = workbook["Dashboard"]
        assert dashboard["A4"].value == "Overall Severity"
        assert dashboard["A6"].value == f"Overall Severity — {severity}"
        assert dashboard["A10"].value == "Final Earnings-Quality Conclusion"
        assert dashboard["A7"].value == (
            f"Final Earnings-Quality Conclusion — {conclusion}"
        )
        analyst = workbook["Analyst Conclusion"]
        assert analyst["A6"].value == "Overall Severity"
        overall = workbook["Overall Severity"]
        assert overall["A1"].value == "Overall Severity"
        assert "Benchmark Severity" not in {
            cell.value for cell in overall[1]
        }
        assert analyst["A7"].value == severity
        assert analyst["E6"].value == "Final Earnings-Quality Conclusion"
        assert analyst["E7"].value == conclusion
    finally:
        workbook.close()


def test_company_ticker_year_filename_and_download_contract(monkeypatch, tmp_path):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    years = (2021, 2022, 2023, 2024, 2025)
    assert run.path == (
        tmp_path / "excel reports"
        / "earnings_quality_analysis_AAPL_2021_2022_2023_2024_2025.xlsx"
    )
    pipeline_result = SimpleNamespace(
        company=APPLE,
        validation=SimpleNamespace(requested_fiscal_years=years),
        excel_output=run.path,
    )
    assert main.downloadable_workbook_path(pipeline_result, APPLE, years) == run.path
    workbook = load_workbook(run.path, data_only=False, read_only=True)
    try:
        assert workbook["Dashboard"]["A2"].value == "Apple Inc. (AAPL)"
        assert workbook["Dashboard"]["A8"].value == "2021, 2022, 2023, 2024, 2025"
        assert workbook["Dashboard"]["B3"].value == "Apple Inc."
        assert workbook["Dashboard"]["E3"].value == "AAPL"
        assert workbook["Dashboard"]["H3"].value == "2021–2025"
        assert workbook["Dashboard"]["L3"].value == 2025
    finally:
        workbook.close()


def test_fresh_five_year_apple_workbook_reopens_and_preserves_fy2024_one_off(
    monkeypatch, tmp_path
):
    run = _finalized_run(
        monkeypatch,
        tmp_path,
        _five_year_apple(),
        APPLE,
        one_off_evidence=_five_year_one_off_fixture(),
    )
    workbook = load_workbook(run.path, data_only=False, read_only=True)
    try:
        normalized = workbook["Normalized Earnings"]
        headers = {cell.value: cell.column for cell in normalized[1]}
        rows = {
            normalized.cell(row, headers["Fiscal Year"]).value: row
            for row in range(2, 7)
        }
        assert set(rows) == {2021, 2022, 2023, 2024, 2025}
        assert normalized.cell(
            rows[2024], headers["One-Off Categories Identified"]
        ).value == "Unusual tax gains/losses"
        assert normalized.cell(
            rows[2023], headers["One-Off Categories Identified"]
        ).value != "Unusual tax gains/losses"
    finally:
        workbook.close()
