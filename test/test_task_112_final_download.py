"""Focused Task 112 tests for the final downloadable analysis."""

import sys
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import main  # noqa: E402
from Data.analysis_validation_gate import (  # noqa: E402
    AnalysisDecision,
    AnalysisValidationGateResult,
    BlockingIssue,
)
from Data.financial_statement_fetcher import FactStatus  # noqa: E402
from Data.manual_corrections import create_manual_correction  # noqa: E402
from main import (  # noqa: E402
    PipelineModules,
    PipelineResult,
    PipelineStageError,
    _apply_correction_overlay,
    _production_excel_output,
    downloadable_workbook_path,
    pipeline_result_matches_selection,
    run_pipeline,
)


COMPANY = SimpleNamespace(company_name="Example Industries", ticker="EXM")
YEARS = (2022, 2023, 2024, 2025)


def _financial_data(*, missing=False, company=COMPANY):
    return pd.DataFrame(
        [
            {
                "Company": company.company_name,
                "Ticker": company.ticker,
                "Fiscal Year": year,
                "Metric": "Revenue",
                "Value": None if missing and year == 2024 else float(year),
                "Units": "USD millions",
                "Source": f"https://www.sec.gov/Archives/{year}",
                "Source Date": f"{year}-12-31",
            }
            for year in YEARS
        ]
    )


def _gate(*, partial=False):
    issues = ()
    if partial:
        issues = (
            BlockingIssue(
                metric="Revenue",
                fiscal_year=2024,
                validation_source="Task 112 fixture",
                blocking_status="MISSING",
                explanation="Revenue evidence is unavailable.",
                manual_correction_used=False,
            ),
        )
    return AnalysisValidationGateResult(
        decision=AnalysisDecision.STOP if partial else AnalysisDecision.CONTINUE,
        can_analyze=not partial,
        requested_fiscal_years=YEARS,
        blocking_issues=issues,
        metric_year_results=(),
        blocking_issue_count=len(issues),
    )


def _interpretations(*, partial=False):
    return {
        "final_conclusion": {
            "earnings_quality_conclusion": (
                "Needs Investigation" if partial else "Existing conclusion"
            ),
            "explanation": (
                "Partial Analysis: FY 2024 Revenue evidence is missing; dependent "
                "outputs remain Unavailable. This conclusion may change if a "
                "verified value is added through the correction panel."
                if partial
                else "Existing complete-analysis conclusion."
            ),
            **({"analysis_status": "Partial Analysis"} if partial else {}),
        }
    }


def _finalized_outputs(corrected_data, *, partial=False):
    summary = pd.DataFrame({
        "Fiscal Year": YEARS,
        "Revenue": (1000.0, 1100.0, 1200.0, 1300.0),
        "Net Income": (100.0, 110.0, 120.0, 130.0),
        "Operating Cash Flow": (150.0, 160.0, 170.0, 180.0),
        "Revenue Growth": (None, 0.10, 0.0909, 0.0833),
        "AR Growth": (None, 0.05, 0.06, 0.07),
        "Depreciation & Amortization": (40.0, 42.0, 44.0, 46.0),
        "Capital Expenditures": (50.0, 52.0, 54.0, 56.0),
        "Stock-Based Compensation": (5.0, 5.5, 6.0, 6.5),
    })
    normalized = pd.DataFrame({
        "Fiscal Year": YEARS,
        "Reported Net Income": summary["Net Income"],
        "Normalized Net Income": (
            (None, None, None, None) if partial else summary["Net Income"]
        ),
        "Normalization Signal": (
            ("Unavailable",) * 4 if partial else ("No Flag",) * 4
        ),
        "Overall Severity": (
            ("Unavailable",) * 4 if partial else ("None",) * 4
        ),
        "Explanation": (
            ("Unavailable",) * 4 if partial else ("No adjustment supplied.",) * 4
        ),
    })
    cash = pd.DataFrame({
        "Fiscal Year": YEARS[1:],
        "AR Growth": (0.05, 0.06, 0.07),
        "Revenue Growth": (0.10, 0.0909, 0.0833),
        "AR Flag": (False, False, False),
        "Explanation": ("Finalized accrual explanation.",) * 3,
    })
    capex = pd.DataFrame({
        "Fiscal Year": YEARS,
        "CapEx": summary["Capital Expenditures"],
        "D&A": summary["Depreciation & Amortization"],
        "CapEx / D&A": (1.25, 1.238095, 1.227273, 1.217391),
        "Exact Rule": ("CapEx / D&A >= 1.20x",) * 4,
        "Status": ("Needs investigation",) * 4,
    })
    empty_analysis = pd.DataFrame({
        "Fiscal Year": YEARS[1:],
        "Explanation": ("Finalized analytical explanation.",) * 3,
    })
    red_flags = pd.DataFrame({
        "Fiscal Year": YEARS[1:],
        "Metric": ("Metric A", "Metric A", "Metric A"),
        "Raw Data": ("Finalized",) * 3,
        "Calculation": ("Finalized",) * 3,
        "Rule": ("Finalized rule",) * 3,
        "Result": ("No Flag",) * 3,
        "Severity": ("None",) * 3,
        "Explanation": ("Finalized red-flag explanation.",) * 3,
    })
    detailed = red_flags[["Fiscal Year", "Metric", "Result", "Severity", "Explanation"]].rename(
        columns={"Severity": "Internal Severity"}
    )
    detailed["Final Severity"] = "Low Risk"
    detailed = detailed[[
        "Fiscal Year", "Metric", "Result", "Internal Severity",
        "Final Severity", "Explanation",
    ]]
    fiscal = pd.DataFrame({
        "Fiscal Year": YEARS[1:],
        "High-Level Severity": ("Low Risk",) * 3,
        "Material Concern Count": (0, 0, 0),
        "Needs Investigation Count": (0, 0, 0),
        "Low Risk Count": (1, 1, 1),
        "Triggered Metrics": ("None",) * 3,
        "Explanation": ("Finalized severity explanation.",) * 3,
    })
    unavailable = pd.DataFrame(columns=[
        "Analysis", "Fiscal Year", "Output", "Status", "Missing Metric",
        "Missing Fiscal Year", "Explanation",
    ])
    if partial:
        unavailable.loc[0] = [
            "Normalized Earnings", 2024, "Normalized Net Income", "Unavailable",
            "One-Off Evidence", 2024, "Validated evidence is unavailable.",
        ]
    calculations = {
        "financial_summary": summary,
        "cash_conversion": cash,
        "normalized_earnings": normalized,
        "excel_support_tables": {"capex": capex},
    }
    analysis = {
        "working_capital": empty_analysis.copy(deep=True),
        "taxes": empty_analysis.copy(deep=True),
        "sbc": empty_analysis.copy(deep=True),
        "red_flags": red_flags,
        "detailed_severity": detailed,
        "fiscal_year_severity": fiscal,
        "overall_severity": pd.DataFrame({
            "Benchmark Severity": ["Existing severity"],
            "Explanation": ["Finalized overall-severity explanation."],
        }),
        "unavailable_outputs": unavailable,
    }
    return calculations, analysis


def _generate(monkeypatch, tmp_path, corrected, *, partial=False, company=COMPANY):
    monkeypatch.setattr(main, "_PROJECT_ROOT", tmp_path)
    structured = (
        {"partial_analysis": {"status": "Partial"}} if partial else {}
    )
    interpretations = _interpretations(partial=partial)
    calculations, analysis = _finalized_outputs(corrected, partial=partial)
    path = _production_excel_output(
        company,
        corrected,
        calculations,
        analysis,
        structured,
        interpretations,
    )
    result = PipelineResult(
        company=company,
        extracted_data=corrected,
        validation=_gate(partial=partial),
        corrected_data=corrected,
        calculations=calculations,
        red_flags_and_severity=analysis,
        structured_ai_input=structured,
        interpretations=interpretations,
        excel_output=path,
    )
    return path, result


def _raw_records(path):
    return pd.read_excel(path, sheet_name="Raw Financial Data")


def test_complete_analysis_produces_a_downloadable_company_year_workbook(
    monkeypatch, tmp_path
):
    path, result = _generate(monkeypatch, tmp_path, _financial_data())

    assert downloadable_workbook_path(result, COMPANY, YEARS) == path
    assert COMPANY.ticker in path.name
    assert all(str(year) in path.name for year in YEARS)
    workbook = load_workbook(path, read_only=True, data_only=False)
    assert workbook["Dashboard"]["A2"].value == "Example Industries (EXM)"
    assert workbook["Dashboard"]["A8"].value == "2022, 2023, 2024, 2025"
    assert workbook["Analyst Conclusion"]["A8"].value == "Existing conclusion"
    workbook.close()


def test_run_pipeline_with_enum_status_returns_downloadable_result(
    monkeypatch, tmp_path
):
    corrected = _financial_data()
    corrected["Status"] = FactStatus.RETRIEVED
    calculations, analysis = _finalized_outputs(corrected)
    modules = PipelineModules(
        identify_company=lambda query: COMPANY,
        extract_data=lambda company, years: corrected,
        validate_data=lambda data, years, corrections: _gate(),
        apply_manual_corrections=lambda data, corrections: data,
        run_calculations=lambda data: calculations,
        run_red_flags_and_severity=lambda calculated: analysis,
        build_structured_ai_input=lambda *args: {},
        run_interpretations=lambda structured: _interpretations(),
        generate_excel=_production_excel_output,
    )
    monkeypatch.setattr(main, "_PROJECT_ROOT", tmp_path)

    result = run_pipeline("EXM", 4, modules=modules)

    assert isinstance(result, PipelineResult)
    assert downloadable_workbook_path(result, COMPANY, YEARS) == result.excel_output
    raw = _raw_records(result.excel_output)
    assert set(raw["Status"]) == {FactStatus.RETRIEVED.value}


def test_partial_analysis_preserves_unavailable_evidence_and_disclosure(
    monkeypatch, tmp_path
):
    path, result = _generate(
        monkeypatch, tmp_path, _financial_data(missing=True), partial=True
    )

    assert downloadable_workbook_path(result, COMPANY, YEARS) == path
    raw = _raw_records(path)
    assert pd.isna(raw.loc[raw["Fiscal Year"].eq(2024), "Value"]).all()
    workbook = load_workbook(path, read_only=True, data_only=False)
    assert workbook["Analyst Conclusion"]["A8"].value == "Needs Investigation"
    assert "partial analysis" in workbook["Analyst Conclusion"]["A14"].value.casefold()
    assert "unavailable" in workbook["Analyst Conclusion"]["A14"].value.casefold()
    workbook.close()


def test_manual_correction_is_refreshed_with_correction_provenance(
    monkeypatch, tmp_path
):
    original = _financial_data()
    record = original.loc[original["Fiscal Year"].eq(2024)].iloc[0].to_dict()
    correction = create_manual_correction(record, 9999.0, "Verified filing correction")
    corrected = _apply_correction_overlay(original, (correction,))

    path, result = _generate(monkeypatch, tmp_path, corrected)
    raw = _raw_records(path)
    corrected_row = raw.loc[raw["Fiscal Year"].eq(2024)].iloc[0]

    assert downloadable_workbook_path(result, COMPANY, YEARS) == path
    assert corrected_row["Value"] == 9999.0
    assert corrected_row["Original Value"] == 2024.0
    assert corrected_row["Correction Reason"] == "Verified filing correction"
    assert corrected_row["Correction Status"] == "ACTIVE MANUAL CORRECTION"


def test_previous_company_or_year_workbook_is_never_served(monkeypatch, tmp_path):
    path, result = _generate(monkeypatch, tmp_path, _financial_data())
    other_company = SimpleNamespace(company_name="Other Corp", ticker="OTHR")

    assert path.is_file()
    assert downloadable_workbook_path(result, other_company, YEARS) is None
    assert downloadable_workbook_path(result, COMPANY, YEARS[:-1]) is None


def test_result_identity_uses_ticker_and_exact_years_not_company_name():
    legal_company = SimpleNamespace(
        company_name="Example Industries, Inc.", ticker="exm"
    )
    result = SimpleNamespace(company=legal_company, validation=_gate())

    assert pipeline_result_matches_selection(result, COMPANY, YEARS)
    assert pipeline_result_matches_selection(result, legal_company, YEARS)
    assert not pipeline_result_matches_selection(
        result,
        SimpleNamespace(company_name=legal_company.company_name, ticker="OTHER"),
        YEARS,
    )
    assert not pipeline_result_matches_selection(result, COMPANY, YEARS[:-1])


def test_workbook_validation_uses_pipeline_company_identity(monkeypatch, tmp_path):
    legal_company = SimpleNamespace(
        company_name="Example Industries, Inc.", ticker="EXM"
    )
    path, result = _generate(
        monkeypatch,
        tmp_path,
        _financial_data(company=legal_company),
        company=legal_company,
    )

    assert downloadable_workbook_path(result, COMPANY, YEARS) == path


def test_stale_workbook_from_another_ticker_is_rejected(monkeypatch, tmp_path):
    other_company = SimpleNamespace(company_name="Other Corp", ticker="OTHR")
    path, result = _generate(
        monkeypatch,
        tmp_path,
        _financial_data(company=other_company),
        company=other_company,
    )

    assert path.is_file()
    assert downloadable_workbook_path(result, COMPANY, YEARS) is None


def test_excel_failure_has_no_result_or_false_download():
    def fail_excel(*_args):
        raise OSError("workbook write failed")

    modules = PipelineModules(
        identify_company=lambda query: COMPANY,
        extract_data=lambda company, years: _financial_data(),
        validate_data=lambda data, years, corrections: _gate(),
        apply_manual_corrections=lambda data, corrections: data,
        run_calculations=lambda data: {"existing calculation": 42},
        run_red_flags_and_severity=lambda calculations: {"existing flag": "Flag"},
        build_structured_ai_input=lambda *args: {},
        run_interpretations=lambda structured: _interpretations(),
        generate_excel=fail_excel,
    )

    with pytest.raises(PipelineStageError) as captured:
        run_pipeline("EXM", 4, modules=modules)

    assert captured.value.failed_stage == "Excel generation"
    assert captured.value.original_message == "workbook write failed"
    assert downloadable_workbook_path(None, COMPANY, YEARS) is None


def test_download_generation_does_not_change_existing_financial_outputs(
    monkeypatch, tmp_path
):
    corrected = _financial_data()
    corrected_before = corrected.copy(deep=True)
    structured = {"company": "EXM"}
    interpretations = _interpretations()
    calculations, analysis = _finalized_outputs(corrected)
    objects_before = deepcopy((calculations, analysis, structured, interpretations))
    monkeypatch.setattr(main, "_PROJECT_ROOT", tmp_path)

    _production_excel_output(
        COMPANY, corrected, calculations, analysis, structured, interpretations
    )

    pd.testing.assert_frame_equal(corrected, corrected_before)
    calculations_before, analysis_before, structured_before, interpretations_before = (
        objects_before
    )
    for key in ("financial_summary", "cash_conversion", "normalized_earnings"):
        pd.testing.assert_frame_equal(calculations[key], calculations_before[key])
    pd.testing.assert_frame_equal(
        calculations["excel_support_tables"]["capex"],
        calculations_before["excel_support_tables"]["capex"],
    )
    for key, frame in analysis.items():
        pd.testing.assert_frame_equal(frame, analysis_before[key])
    assert structured == structured_before
    assert interpretations == interpretations_before


if __name__ == "__main__":
    exit_code = pytest.main([__file__, "-q"])
    if exit_code == 0:
        print("Task 112 final download tests passed")
    raise SystemExit(exit_code)
